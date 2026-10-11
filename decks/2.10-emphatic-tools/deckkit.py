"""ArabicHub deck toolkit: RTL-aware slide primitives + build animations.

Slides are 960 x 540 pt (16:9). All coordinates below are in points.
Text blocks use a small line syntax (see `paras_from`):
    "~text"   transliteration line (italic, muted)
    "#text"   card/section heading (16 pt, bold)
    "!text"   small note line (11 pt, muted)
    other     Arabic if its first strong letter is Arabic, else English
Inline markup: {b:..} blue tool, {g:..} green, {r:..} red, {w:..} wine,
{k:..} bold, {i:..} italic, {u:..} white bold.
"""
import re
import unicodedata
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- palette
WINE = "8E2453"
TINT = "F7E6EE"
TINT2 = "FBF3F7"
BLUE = "33618F"
LBLUE = "E8EFF7"
GREEN = "2E7D32"
LGREEN = "E7F4E8"
RED = "C62828"
LRED = "FBEAEA"
GREYBOX = "A7A7A7"
TITLEFILL = "F2F2F2"
INK = "262626"
MUTED = "5F5F5F"
LINE = "D9CDD3"
WHITE = "FFFFFF"
GOLD = "C9A227"
ROAD = "4A4A4A"
FLAGGREEN = "2E8B3E"
INACTIVE = "D8D8D8"

EN_FONT = "Chalkboard SE"   # course English face (Keynote / macOS)
AR_FONT = "Geeza Pro"       # course Arabic face (Keynote / macOS)

AR_SIZE = 19
EN_SIZE = 12

RLM = "‏"
LRM = "‎"

AR_RE = re.compile(r"[؀-ۿݐ-ݿﭐ-﷿ﹰ-﻿]")
LAT_RE = re.compile(r"[A-Za-zÀ-ɏḀ-ỿ]")
MARKS = set(chr(c) for c in range(0x064B, 0x0653)) | {"ٰ"}
SHADDA = "ّ"
MARKUP = re.compile(r"\{([bgrwkiut]):(.*?)\}")

STYLE = {
    "b": dict(color=BLUE, bold=True),
    "g": dict(color=GREEN, bold=True),
    "r": dict(color=RED, bold=True),
    "w": dict(color=WINE, bold=True),
    "k": dict(bold=True),
    "i": dict(italic=True),
    "u": dict(color=WHITE, bold=True),
    "t": dict(color=MUTED, italic=True),
}


def is_ar(ch):
    return bool(AR_RE.match(ch))


def first_strong(text):
    for ch in MARKUP.sub(lambda m: m.group(2), text):
        if is_ar(ch):
            return "ar"
        if LAT_RE.match(ch):
            return "en"
    return "en"


def strip_final_marks(text):
    """Course convention: no vowel sign on the final written letter of a
    word.  A final shadda is kept (it is part of the consonant: كُلّ);
    any vowel/tanwin/sukun on the final letter is removed.  Inline markup
    tokens such as {b: and } are transparent, so a word split across
    markup runs keeps its internal vowels."""
    out = []
    i = 0
    n = len(text)
    while i < n:
        ch = text[i]
        if ch in MARKS:
            j = i
            while j < n and text[j] in MARKS:
                j += 1
            cluster = text[i:j]
            k = j
            while k < n:
                if text[k] == "}":
                    k += 1
                elif re.match(r"\{[a-z]:", text[k:k + 3]):
                    k += 3
                else:
                    break
            nxt = text[k] if k < n else " "
            # the word continues only if a letter (or tatweel) follows
            if unicodedata.category(nxt) not in ("Lo", "Lm"):
                cluster = SHADDA if SHADDA in cluster else ""
            out.append(cluster)
            i = j
            continue
        out.append(ch)
        i += 1
    return "".join(out)


def ar_punct(text):
    return text.replace("?", "؟").replace(",", "،").replace(";", "؛")


# ---------------------------------------------------------------- text
class Para:
    def __init__(self, text, size=None, color=INK, bold=False, italic=False,
                 align=None, kind=None, space_after=2, line=1.08):
        self.text = text
        self.kind = kind or first_strong(text)
        self.size = size
        self.color = color
        self.bold = bold
        self.italic = italic
        self.align = align
        self.space_after = space_after
        self.line = line


def paras_from(spec, en_size=EN_SIZE, ar_size=AR_SIZE, color=INK,
               align="auto", head_color=WINE):
    if isinstance(spec, Para):
        if align != "auto" and spec.align is None:
            spec.align = align
        return [spec]
    if isinstance(spec, (list, tuple)):
        out = []
        for s in spec:
            out += paras_from(s, en_size, ar_size, color, align, head_color)
        return out
    out = []
    for line in str(spec).split("\n"):
        if line.startswith("~"):
            p = Para(line[1:], size=en_size, color=MUTED, italic=True, kind="en")
        elif line.startswith("#"):
            t = line[1:]
            p = Para(t, size=16, color=head_color, bold=True)
            if p.kind == "ar":
                p.size = max(ar_size, 18)
        elif line.startswith("!"):
            p = Para(line[1:], size=max(en_size - 1, 10), color=MUTED, kind="en")
        else:
            p = Para(line, color=color)
            p.size = ar_size if p.kind == "ar" else en_size
        if align != "auto":
            p.align = align
        out.append(p)
    return out


def _rpr(r, size, color, bold, italic, arabic):
    rPr = etree.SubElement(r, qn("a:rPr"))
    rPr.set("lang", "ar-SA" if arabic else "en-GB")
    if arabic:
        rPr.set("altLang", "en-GB")
    rPr.set("sz", str(int(round(size * 100))))
    rPr.set("b", "1" if bold else "0")
    rPr.set("i", "1" if italic else "0")
    rPr.set("dirty", "0")
    sf = etree.SubElement(rPr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
    font = AR_FONT if arabic else EN_FONT
    etree.SubElement(rPr, qn("a:latin")).set("typeface", font)
    etree.SubElement(rPr, qn("a:cs")).set("typeface", AR_FONT)
    return rPr


def _split_script(text):
    """Split a plain string into (segment, is_arabic) chunks so Arabic
    words inside an English line get the Arabic face and size."""
    chunks = []
    cur = ""
    cur_ar = None
    for ch in text:
        if is_ar(ch) or ch in MARKS or ch == "ـ":
            a = True
        elif LAT_RE.match(ch) or ch.isdigit():
            a = False
        else:
            a = cur_ar  # neutral: stick with current run
        if cur_ar is None:
            cur_ar = a
        if a != cur_ar and cur:
            chunks.append((cur, bool(cur_ar)))
            cur = ""
        cur_ar = a if a is not None else cur_ar
        cur += ch
    if cur:
        chunks.append((cur, bool(cur_ar)))
    return chunks


def write_paras(tf, paras, anchor="t", margins=(6, 6, 6, 4), autofit=False):
    tf.word_wrap = True
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = \
        [Pt(m) for m in margins]
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[anchor]
    txBody = tf._txBody
    bodyPr = txBody.find(qn("a:bodyPr"))
    for child in list(bodyPr):
        bodyPr.remove(child)
    etree.SubElement(bodyPr, qn("a:noAutofit"))
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    for para in paras:
        p = etree.SubElement(txBody, qn("a:p"))
        pPr = etree.SubElement(p, qn("a:pPr"))
        rtl = para.kind == "ar"  # noqa
        align = para.align or ("r" if rtl else "l")
        pPr.set("algn", {"l": "l", "r": "r", "c": "ctr", "center": "ctr"}[align])
        pPr.set("rtl", "1" if rtl else "0")
        ln = etree.SubElement(pPr, qn("a:lnSpc"))
        etree.SubElement(ln, qn("a:spcPct")).set("val", str(int(para.line * 100000)))
        sa = etree.SubElement(pPr, qn("a:spcAft"))
        etree.SubElement(sa, qn("a:spcPts")).set("val", str(int(para.space_after * 100)))
        etree.SubElement(pPr, qn("a:buNone"))
        text = strip_final_marks(para.text)
        if not rtl:
            # keep ': , ;' after an Arabic word on the English (LTR) side
            text = re.sub(r"([\u0621-\u064A\u0640][\u064B-\u0652\u0670]*\}?)([:;,])",
                          "\\1" + LRM + "\\2", text)
        has_ar = bool(AR_RE.search(text))
        # directional mark at paragraph start keeps Keynote's direction
        lead = RLM if rtl else (LRM if has_ar else "")
        pieces = []
        pos = 0
        for m in MARKUP.finditer(text):
            if m.start() > pos:
                pieces.append((text[pos:m.start()], {}))
            pieces.append((m.group(2), STYLE[m.group(1)]))
            pos = m.end()
        if pos < len(text):
            pieces.append((text[pos:], {}))
        first = True
        for seg, st in pieces:
            for chunk, arab in _split_script(seg):
                if rtl:
                    chunk = ar_punct(chunk)
                if first:
                    chunk = lead + chunk
                    first = False
                r = etree.SubElement(p, qn("a:r"))
                size = para.size
                if arab and not rtl:
                    size = size + 3 if size < 18 else size
                _rpr(r, size, st.get("color", para.color),
                     st.get("bold", para.bold), st.get("italic", para.italic),
                     arab)
                etree.SubElement(r, qn("a:t")).text = chunk
        if first:  # empty paragraph
            end = etree.SubElement(p, qn("a:endParaRPr"))
            end.set("sz", str(int((para.size or EN_SIZE) * 100)))


# ---------------------------------------------------------------- shapes
def _clean(sh):
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)


def shape(c, kind, x, y, w, h, fill=None, line=None, lw=1.0, radius=None,
          rot=0):
    sh = c.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
    _clean(sh)
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / max(1, min(w, h)))
    if rot:
        sh.rotation = rot
    return sh


def tb(c, x, y, w, h, spec, fill=None, line=None, lw=1.0, radius=None,
       anchor="t", margins=(7, 7, 7, 4), en_size=EN_SIZE, ar_size=AR_SIZE,
       color=INK, align="auto", rot=0, head_color=WINE):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = shape(c, kind, x, y, w, h, fill, line, lw, radius, rot)
    paras = paras_from(spec, en_size, ar_size, color, align, head_color)
    write_paras(sh.text_frame, paras, anchor, margins)
    return sh


def group(c):
    return c.add_group_shape()


def left_arrow(c, x, y, w, h, fill=WINE):
    return shape(c, MSO_SHAPE.LEFT_ARROW, x, y, w, h, fill)


# ---------------------------------------------------------------- animation
class Builds:
    """Ordered build steps for a slide.

    clicks: list of click groups; each click is a list of waves; each wave
    is a list of (spid, effect).  Wave 0 starts on the click; each later
    wave starts automatically after the previous one (used for table cells
    revealed one by one, right to left)."""

    def __init__(self):
        self.clicks = []

    def click(self, shapes, eff="fade"):
        shapes = shapes if isinstance(shapes, (list, tuple)) else [shapes]
        self.clicks.append([[(s.shape_id, eff) for s in shapes]])

    def chain(self, shapes, eff="fade"):
        """One click, then each shape follows the previous automatically."""
        if not shapes:
            return
        self.clicks.append([[(s.shape_id, eff)] for s in shapes])

    def with_(self, shapes, eff="fade"):
        shapes = shapes if isinstance(shapes, (list, tuple)) else [shapes]
        self.clicks[-1][-1] += [(s.shape_id, eff) for s in shapes]


DUR = {"fade": 400, "wipe": 500}


def timing_xml(builds):
    if not builds.clicks:
        return None
    ids = iter(range(1, 100000))
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    nid = lambda: str(next(ids))

    def el(parent, tag, **attrs):
        e = etree.SubElement(parent, qn(tag))
        for k, v in attrs.items():
            e.set(k, v)
        return e

    timing = etree.Element(qn("p:timing"), nsmap={"p": P})
    tnLst = el(timing, "p:tnLst")
    root_par = el(tnLst, "p:par")
    root = el(root_par, "p:cTn", id=nid(), dur="indefinite", restart="never",
              nodeType="tmRoot")
    rc = el(root, "p:childTnLst")
    seq = el(rc, "p:seq", concurrent="1", nextAc="seek")
    main = el(seq, "p:cTn", id=nid(), dur="indefinite", nodeType="mainSeq")
    mc = el(main, "p:childTnLst")
    grp = 0
    for click in builds.clicks:
        cpar = el(mc, "p:par")
        ctn = el(cpar, "p:cTn", id=nid(), fill="hold")
        el(el(ctn, "p:stCondLst"), "p:cond", delay="indefinite")
        cc = el(ctn, "p:childTnLst")
        t = 0
        for wi, wave in enumerate(click):
            wpar = el(cc, "p:par")
            wtn = el(wpar, "p:cTn", id=nid(), fill="hold")
            el(el(wtn, "p:stCondLst"), "p:cond", delay=str(t))
            wc = el(wtn, "p:childTnLst")
            longest = 0
            for ei, (spid, eff) in enumerate(wave):
                node = "clickEffect" if (wi == 0 and ei == 0) else (
                    "afterEffect" if ei == 0 else "withEffect")
                preset, sub, filt = ("10", "0", "fade") if eff == "fade" else \
                    ("22", "2", "wipe(right)")
                epar = el(wc, "p:par")
                etn = el(epar, "p:cTn", id=nid(), presetID=preset,
                         presetClass="entr", presetSubtype=sub, fill="hold",
                         grpId=str(grp), nodeType=node)
                el(el(etn, "p:stCondLst"), "p:cond", delay="0")
                ec = el(etn, "p:childTnLst")
                st = el(ec, "p:set")
                cb = el(st, "p:cBhvr")
                stn = el(cb, "p:cTn", id=nid(), dur="1", fill="hold")
                el(el(stn, "p:stCondLst"), "p:cond", delay="0")
                el(el(cb, "p:tgtEl"), "p:spTgt", spid=str(spid))
                el(el(cb, "p:attrNameLst"), "p:attrName").text = "style.visibility"
                el(el(st, "p:to"), "p:strVal", val="visible")
                ae = el(ec, "p:animEffect", transition="in", filter=filt)
                acb = el(ae, "p:cBhvr")
                el(acb, "p:cTn", id=nid(), dur=str(DUR[eff]))
                el(el(acb, "p:tgtEl"), "p:spTgt", spid=str(spid))
                longest = max(longest, DUR[eff])
            t += longest if wi < len(click) - 1 else 0
    prev = el(seq, "p:prevCondLst")
    c1 = el(prev, "p:cond", evt="onPrev", delay="0")
    el(el(c1, "p:tgtEl"), "p:sldTgt")
    nxt = el(seq, "p:nextCondLst")
    c2 = el(nxt, "p:cond", evt="onNext", delay="0")
    el(el(c2, "p:tgtEl"), "p:sldTgt")
    return timing


def attach_timing(slide, builds):
    t = timing_xml(builds)
    if t is None:
        return
    sld = slide._element
    old = sld.find(qn("p:timing"))
    if old is not None:
        sld.remove(old)
    ext = sld.find(qn("p:extLst"))
    if ext is not None:
        ext.addprevious(t)
    else:
        sld.append(t)


def new_presentation():
    prs = Presentation()
    prs.slide_width = Pt(960)
    prs.slide_height = Pt(540)
    return prs
