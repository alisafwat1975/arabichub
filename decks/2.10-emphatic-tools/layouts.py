"""Course frame and reusable slide layouts for ArabicHub Chapter 2.10."""
from pptx.enum.shapes import MSO_SHAPE
from deckkit import *  # noqa: F401,F403

CHAPTER_NO = "2.10"
CHAPTER_EN = "Emphatic Tools"
CHAPTER_AR = "أَدَوَات التَّوْكِيد"
CHAPTER_TR = "adawāt at-tawkīd"
TEACHER = "Mrs Layla"
# Chapter group shown in the sidebar (as in the course chapter-list view).
CHAPTER_LIST = [("2.10", "Emphatic Tools", True),
                ("2.11", "Vocative Tools", False),
                ("2.12", "Exclamation Tools", False)]

X0, X1 = 96, 940          # content area
CW = X1 - X0
STOP_NAMES = ["What Emphasis Means", "Inside a Sentence",
              "Before & Within a Verb", "Restriction & Completeness",
              "Putting It Together", "Clearing the Confusion"]


def ex(ar, tr, en):
    """Arabic example + one gloss line (transliteration — English)."""
    return [ar, Para("{t:%s}  —  %s" % (tr, en), size=EN_SIZE, kind="en")]


# ---------------------------------------------------------------- frame
def sidebar(c):
    circ = shape(c, MSO_SHAPE.OVAL, 11, 8, 58, 58, "111111")
    write_paras(circ.text_frame, [Para(CHAPTER_NO, size=17, color=WHITE,
                                       bold=True, align="c")], "m", (0, 0, 0, 0))
    y = 76
    heights = [210, 112, 112]
    for (no, name, active), h in zip(CHAPTER_LIST, heights):
        w = 44 if active else 36
        cx, cy = 40, y + h / 2
        fill = WINE if active else INACTIVE
        if active:
            paras = [Para(name, size=19, color=WHITE, bold=True, align="c")]
        else:
            paras = [Para(no, size=9.5, color="444444", bold=True, align="c",
                          space_after=0, line=1.0),
                     Para(name, size=9.5, color="444444", align="c",
                          space_after=0, line=1.0)]
        r = tb(c, cx - h / 2, cy - w / 2, h, w, paras, fill=fill, anchor="m",
               margins=(2, 0, 2, 0), rot=270)
        y += h + 6
    return r


def frame(s, title, stop=None, page=None, footer_mid=None):
    c = s.shapes
    sidebar(c)
    tb(c, X0, 12, CW, 44, Para(title, size=21, color=WINE, bold=True),
       fill=TITLEFILL, line=GREYBOX, lw=1.25, radius=10, anchor="m",
       margins=(14, 2, 120 if stop else 14, 2))
    if stop:
        tb(c, X1 - 112, 21, 102, 26, Para(f"Stop {stop} of 6", size=11,
           color=WHITE, bold=True, align="c"), fill=WINE, radius=13,
           anchor="m", margins=(2, 0, 2, 0))
    shape(c, MSO_SHAPE.RECTANGLE, X0, 506, CW, 1.2, WINE)
    tb(c, X0, 509, 200, 26, Para(TEACHER, size=14, color=WINE, bold=True),
       anchor="m", margins=(0, 0, 0, 0))
    mid = footer_mid or f"Chapter {CHAPTER_NO} · {CHAPTER_EN} · {CHAPTER_AR}"
    tb(c, X0 + 200, 509, CW - 260, 26, Para(mid, size=14, color=MUTED,
       align="c"), anchor="m", margins=(0, 0, 0, 0))
    if page:
        tb(c, X1 - 60, 509, 60, 26, Para(str(page), size=14, color=MUTED,
           align="r"), anchor="m", margins=(0, 0, 0, 0))


def banner(s, B, text, y=62, h=32, build=False):
    g = group(s.shapes)
    shape(g.shapes, MSO_SHAPE.RECTANGLE, X0, y, CW, h, TINT)
    shape(g.shapes, MSO_SHAPE.RECTANGLE, X0, y, 6, h, WINE)
    tb(g.shapes, X0 + 12, y, CW - 18, h, Para(text, size=16, color=INK),
       anchor="m", margins=(4, 1, 4, 1))
    if build:
        B.click(g)
    return g


def note(c, x, y, w, h, text, label="NOTE", color=BLUE, fill=LBLUE, size=16):
    g = group(c)
    shape(g.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, color, 1,
          radius=8)
    paras = [Para(label, size=11, color=color, bold=True)]
    paras += paras_from(text, en_size=size, ar_size=max(size, 18))
    t = tb(g.shapes, x + 4, y + 2, w - 8, h - 4, paras, anchor="m",
           margins=(8, 4, 8, 4))
    return g


def card(c, x, y, w, h, title, body, color=WINE, body_fill=TINT2, title_h=30,
         en_size=EN_SIZE, ar_size=18, title_size=16, anchor="t", align="auto"):
    g = group(c)
    shape(g.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, body_fill, color,
          1.2, radius=10)
    # header: rounded top + square lower corners, then the title text on top
    shape(g.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, title_h, color,
          radius=10)
    shape(g.shapes, MSO_SHAPE.RECTANGLE, x, y + title_h - 10, w, 10, color)
    tb(g.shapes, x, y, w, title_h, Para(title, size=title_size, color=WHITE,
       bold=True, align="c"), anchor="m", margins=(6, 1, 6, 1))
    tb(g.shapes, x + 2, y + title_h + 2, w - 4, h - title_h - 4, body,
       anchor=anchor, en_size=en_size, ar_size=ar_size, align=align,
       margins=(8, 6, 8, 4))
    return g


def exbox(c, x, y, w, h, ar, tr, en, fill=WHITE, line=LINE, align="r",
          ar_size=AR_SIZE, label=None, label_color=WINE):
    spec = []
    if label:
        spec.append(Para(label, size=11, color=label_color, bold=True,
                         align=align, kind="en"))
    spec.append(Para(ar, size=ar_size, align=align, kind="ar"))
    spec.append(Para("{t:%s}  —  %s" % (tr, en), size=EN_SIZE, align=align,
                     kind="en"))
    return tb(c, x, y, w, h, spec, fill=fill, line=line, radius=8,
              anchor="m", margins=(10, 6, 10, 4))


# ---------------------------------------------------------------- tables
def table(s, B, x, y, colw, header, rows, rowh=44, header_h=28,
          ar_size=18, en_size=11.5, fills=None, align="c", build=True,
          header_fill=WINE):
    """colw / header / row cells are given in READING order: the first
    column is drawn at the RIGHT.  A cell is a spec string or (spec, fill)."""
    c = s.shapes
    xs = []
    cx = x + sum(colw)
    for w in colw:
        cx -= w
        xs.append(cx)
    hdr = []
    for i, (hx, w) in enumerate(zip(xs, colw)):
        hdr.append(tb(c, hx, y, w, header_h, Para(header[i], size=12,
                   color=WHITE, bold=True, align="c"), fill=header_fill,
                   line=WHITE, lw=1.5, anchor="m", margins=(4, 1, 4, 1)))
    if build:
        B.chain(hdr)
    yy = y + header_h
    grid = []
    rh = rowh if isinstance(rowh, (list, tuple)) else [rowh] * len(rows)
    for ri, row in enumerate(rows):
        cells = []
        for i, (cx, w) in enumerate(zip(xs, colw)):
            cell = row[i]
            fill = (WHITE if ri % 2 == 0 else TINT2)
            if isinstance(cell, tuple):
                cell, fill = cell
            if fills and fills.get((ri, i)):
                fill = fills[(ri, i)]
            cells.append(tb(c, cx, yy, w, rh[ri], cell, fill=fill, line=LINE,
                            lw=0.75, anchor="m", ar_size=ar_size,
                            en_size=en_size, align=align,
                            margins=(6, 4, 6, 3)))
        if build:
            B.chain(cells)
        grid.append(cells)
        yy += rh[ri]
    return grid


# ---------------------------------------------------------------- drills
def drill(s, B, items, y0=104, rowh=62, gap=6, widths=(30, 372, 196, 226),
          ar_size=18, show_numbers=True):
    """Each item: dict(q=prompt spec, cue=cue spec, a=answer spec).
    Reading order (from the right): number → prompt → cue → answer.
    Builds: prompt (wipe from right) → cue → answer."""
    c = s.shapes
    gaps = (CW - sum(widths)) / (len(widths) - 1)
    xs = []
    cx = X1
    for w in widths:
        cx -= w
        xs.append(cx)
        cx -= gaps
    y = y0
    for n, it in enumerate(items, 1):
        g = group(c)
        circ = shape(g.shapes, MSO_SHAPE.OVAL, xs[0] + 2, y + rowh / 2 - 13, 26,
                     26, WINE)
        write_paras(circ.text_frame, [Para(str(n), size=12, color=WHITE,
                    bold=True, align="c")], "m", (0, 0, 0, 0))
        tb(g.shapes, xs[1], y, widths[1], rowh, it["q"], fill=WHITE, line=LINE,
           radius=8, anchor="m", ar_size=ar_size, align="r")
        cue = tb(c, xs[2], y, widths[2], rowh, it["cue"], fill=LBLUE,
                 line=BLUE, lw=0.75, radius=8, anchor="m", color=BLUE,
                 ar_size=ar_size, align="c", en_size=11.5)
        ans = tb(c, xs[3], y, widths[3], rowh, it["a"], fill=LGREEN,
                 line=GREEN, lw=1, radius=8, anchor="m", color=INK,
                 ar_size=ar_size, align="c", en_size=11.5)
        B.click(g, "wipe")
        B.click(cue)
        B.click(ans)
        y += rowh + gap
    return y


def drill_header(s, y=74, labels=("Prompt", "Cue", "Answer"),
                 widths=(30, 372, 196, 226)):
    c = s.shapes
    gaps = (CW - sum(widths)) / (len(widths) - 1)
    cx = X1 - widths[0] - gaps
    for lab, w in zip(labels, widths[1:]):
        cx -= w
        tb(c, cx, y, w, 22, Para(lab.upper(), size=11, color=MUTED, bold=True,
           align="c"), anchor="m", margins=(0, 0, 0, 0))
        cx -= gaps


def fixrow(s, B, y, wrong, right, why, h=66, ar_size=18):
    """✗ on the right → arrow left → ✓ → reason (leftmost)."""
    c = s.shapes
    w_wrong, w_right, w_why = 286, 286, 214
    xw = X1 - w_wrong
    xa = xw - 40
    xr = xa - 6 - w_right
    xy = X0
    gw = tb(c, xw, y, w_wrong, h, [Para("✗ Incorrect", size=11, color=RED,
            bold=True, align="r", kind="en"), Para(wrong, size=ar_size,
            color=INK, align="r", kind="ar")], fill=LRED, line=RED, lw=1,
            radius=8, anchor="m", margins=(8, 3, 8, 3))
    ar = left_arrow(c, xa + 4, y + h / 2 - 11, 30, 22, WINE)
    right_spec = [Para("✓ Correct", size=11, color=GREEN, bold=True,
                       align="r", kind="en")]
    right_spec += paras_from(right, ar_size=ar_size, align="r")
    gr = tb(c, xr, y, w_right, h, right_spec, fill=LGREEN, line=GREEN, lw=1,
            radius=8, anchor="m", margins=(8, 3, 8, 3))
    gy = tb(c, xy, y, w_why, h, why, fill=None, anchor="m", en_size=11.5,
            ar_size=16, margins=(4, 2, 4, 2))
    B.click(gw, "wipe")
    B.click([ar, gr])
    B.with_(gy)


def bubble(s, B, y, side, name, ar, tr, en, h=78, w=560):
    c = s.shapes
    x = X1 - w if side == "r" else X0
    fill = TINT if side == "r" else LBLUE
    line = WINE if side == "r" else BLUE
    g = group(c)
    tb(g.shapes, x, y, w, h, [Para(ar, size=AR_SIZE, align="r", kind="ar"),
       Para("{t:%s}" % tr, size=EN_SIZE, align="r", kind="en"),
       Para(en, size=EN_SIZE, align="r", kind="en")], fill=fill, line=line,
       radius=14, anchor="m", margins=(14, 6, 14, 4))
    nx = x + w + 6 if side == "r" else x + w + 6
    if side == "r":
        tb(g.shapes, x - 92, y + h / 2 - 14, 86, 28, Para(name, size=12,
           color=WINE, bold=True, align="r"), anchor="m", margins=(0, 0, 0, 0))
    else:
        tb(g.shapes, x + w + 6, y + h / 2 - 14, 86, 28, Para(name, size=12,
           color=BLUE, bold=True, align="l"), anchor="m", margins=(0, 0, 0, 0))
    B.click(g, "wipe")
    return g


def legend(s, items, y=482):
    """Explain any colour coding used on the slide."""
    c = s.shapes
    x = X0
    for col, fill, label in items:
        shape(c, MSO_SHAPE.ROUNDED_RECTANGLE, x, y + 5, 14, 12, fill, col, 1,
              radius=3)
        t = tb(c, x + 18, y, 230, 22, Para(label, size=10.5, color=MUTED),
               anchor="m", margins=(0, 0, 0, 0))
        x += 22 + 8 * len(label) * 0.62 + 30
