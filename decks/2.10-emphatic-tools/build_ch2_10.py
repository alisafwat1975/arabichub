"""Build ArabicHub Chapter 2.10 — Emphatic Tools (أَدَوَات التَّوْكِيد).

Run:  python3 build_ch2_10.py   (writes out/Chapter-2.10-Emphatic-Tools.pptx)
The .pptx is the programmatic build source; open/convert it in Apple
Keynote with convert_to_keynote.sh to produce the native .key and the
Keynote-exported PDF.
"""
import os
import sys
from pptx.enum.shapes import MSO_SHAPE
from layouts import *  # noqa: F401,F403

SLIDES = []


def slide(part):
    def deco(fn):
        SLIDES.append((part, fn))
        return fn
    return deco


class Ctx:
    ranges = {}
    page = 0


def rng(ctx, part):
    a, b = ctx.ranges[part]
    return f"slides {a}–{b}"


# =================================================================== OPENING
@slide("open")
def s_title(s, B, ctx):
    c = s.shapes
    sidebar(c)
    tb(c, X0, 40, CW, 30, Para("Chapter 2.10", size=18, color=MUTED,
       bold=True, align="c"))
    tb(c, X0, 72, CW, 74, Para("Emphatic Tools", size=54, color=WINE,
       bold=True, align="c"), anchor="m")
    tb(c, X0, 150, CW, 70, Para("أَدَوَات التَّوْكِيد", size=40, color=INK,
       align="c", kind="ar"), anchor="m", margins=(4, 10, 4, 4))
    tb(c, X0, 222, CW, 26, Para("{t:adawāt at-tawkīd}  ·  tools of emphasis",
       size=16, align="c"), anchor="m")
    tb(c, X0 + 80, 258, CW - 160, 50, Para(
        "How Arabic adds force, certainty, limits and totality to what you say",
        size=18, color=INK, align="c"), anchor="m")
    tools = [("إِنّ", "inna"), ("أَنّ", "anna"), ("قَد", "qad"),
             ("لَقَد", "laqad"), ("لَـ", "la-"), ("إِنَّمَا", "innamā"),
             ("نَفْس", "nafs"), ("عَيْن", "ʿayn"), ("كُلّ", "kull"),
             ("جَمِيع", "jamīʿ")]
    w, gap = 76, 8
    x = X1 - (CW - (10 * w + 9 * gap)) / 2 - w
    chips = []
    for ar, tr in tools:  # first tool at the RIGHT
        chips.append(tb(c, x, 330, w, 70, [Para(ar, size=20, align="c",
                     kind="ar"), Para("{t:%s}" % tr, size=11, align="c",
                     kind="en")], fill=TINT, line=WINE, lw=1, radius=10,
                     anchor="m", margins=(2, 8, 2, 3)))
        x -= w + gap
    B.chain(chips)
    tb(c, X0, 424, CW, 30, Para("The ten tools of this chapter, read from the right",
       size=12, color=MUTED, align="c"))
    shape(c, MSO_SHAPE.RECTANGLE, X0 + 320, 466, CW - 640, 1.2, WINE)
    tb(c, X0, 472, CW, 30, Para("with Mrs Layla", size=18, color=WINE,
       bold=True, align="c"), anchor="m")


@slide("open")
def s_bigpicture(s, B, ctx):
    frame(s, "The Big Picture: Ways Arabic Adds Emphasis", page=ctx.page)
    banner(s, B, "Arabic has different kinds of emphatic tools — particles, an "
           "attached prefix, and ordinary nouns.")
    cols = [
        ("Particles: on a Sentence or Verb", [
            ("1", "إِنّ", "inna", "indeed, truly"),
            ("2", "أَنّ", "anna", "that"),
            ("3", "قَد", "qad", "already / may"),
            ("4", "لَقَد", "laqad", "indeed (+ past verb)"),
            ("6", "إِنَّمَا", "innamā", "only, merely")]),
        ("An Attached Prefix", [
            ("5", "لَـ", "la-", "truly — joined to a word"),
            ("", "لَقَد = لَـ + قَد", "laqad = la- + qad",
             "Tool 4 contains it")]),
        ("Nouns Used for Emphasis", [
            ("7", "نَفْس", "nafs", "self, same"),
            ("8", "عَيْن", "ʿayn", "the very same"),
            ("9", "كُلّ", "kull", "all, every, whole"),
            ("10", "جَمِيع", "jamīʿ", "all, altogether")]),
    ]
    w, gap = 272, 14
    x = X1 - w
    for title, items in cols:
        body = []
        for n, ar, tr, en in items:
            body.append(Para(ar, size=19, align="r", kind="ar", space_after=0))
            lead = f"{{w:Tool {n}}} · " if n else ""
            body.append(Para(f"{lead}{{t:{tr}}} — {en}", size=12, align="r",
                             kind="en", space_after=7))
        g = card(s.shapes, x, 104, w, 336, title, body, title_size=15)
        B.click(g, "wipe")
        x -= w + gap
    st = tb(s.shapes, X0, 450, CW, 44, Para(
        "{w:And with no tool at all:} repetition · intonation · context",
        size=16, align="c"), fill=LBLUE, line=BLUE, radius=10, anchor="m")
    B.click(st)


@slide("open")
def s_whatdoes(s, B, ctx):
    frame(s, "What Does Emphasis Do?", page=ctx.page)
    banner(s, B, "Same information — but the speaker now presents it as certain.")
    a = exbox(s.shapes, X1 - 390, 108, 390, 92, "الْبَيْت كَبِير",
              "al-baytu kabīrun", "The house is big.", label="NEUTRAL STATEMENT")
    arw = left_arrow(s.shapes, X0 + 400, 140, 44, 30)
    b = exbox(s.shapes, X0, 108, 390, 92, "{b:إِنّ} الْبَيْت كَبِير",
              "inna l-bayta kabīrun", "The house is indeed big.",
              fill=LBLUE, line=BLUE, label="EMPHASIZED STATEMENT",
              label_color=BLUE)
    B.click(a, "wipe")
    B.click([arw, b], "wipe")
    w, gap = 272, 14
    x = X1 - w
    items = [
        ("The Facts Stay the Same", "Both sentences say the house is big. "
         "Emphasis adds no new information."),
        ("The Stance Changes", "With {b:إِنّ} the speaker insists it is true — "
         "useful when the listener doubts it or the point matters."),
        ("No Word-for-Word Match", "English uses voice stress, 'really', "
         "'indeed' or 'do'. Arabic often uses grammar. Translate the effect, "
         "not the word."),
    ]
    for t, body in items:
        g = card(s.shapes, x, 220, w, 170, t, Para(body, size=14), title_size=15,
                 anchor="m")
        B.click(g)
        x -= w + gap
    n = note(s.shapes, X0, 402, CW, 66, "In this chapter, {b:blue} always marks "
             "the emphatic tool in an example.", label="COLOUR KEY", size=14)
    B.click(n)


@slide("open")
def s_three_strengths(s, B, ctx):
    frame(s, "Three Strengths of a Statement", page=ctx.page)
    banner(s, B, "A teaching model: three ways to present the same situation.")
    data = [
        ("1 · Neutral Information", "الْمَطْعَم مَفْتُوح", "al-maṭʿamu maftūḥun",
         "The restaurant is open.", "Simple information. No special stress."),
        ("2 · Reinforced Assertion", "{b:إِنّ} الْمَطْعَم مَفْتُوح",
         "inna l-maṭʿama maftūḥun", "The restaurant is definitely open.",
         "The speaker insists: it is open (perhaps someone doubted it)."),
        ("3 · Restricted Meaning", "{b:إِنَّمَا} يَفْتَح الْمَطْعَم مَسَاءً",
         "innamā yaftaḥu l-maṭʿamu masāʾan",
         "The restaurant opens only in the evening.",
         "The statement is limited to one case — and excludes the others."),
    ]
    w, aw = 256, 38
    x = X1 - w
    for i, (t, ar, tr, en, ex_) in enumerate(data):
        body = [Para(ar, size=19, align="c", kind="ar"),
                Para("{t:%s}" % tr, size=12, align="c", kind="en"),
                Para(en, size=12, align="c", kind="en", bold=True,
                     space_after=8),
                Para(ex_, size=12, align="c", kind="en", color=MUTED)]
        g = card(s.shapes, x, 106, w, 236, t, body, title_size=15,
                 anchor="m")
        B.click(g, "wipe")
        if i < 2:
            a = left_arrow(s.shapes, x - aw + 4, 206, aw - 8, 28)
            B.click(a)
        x -= w + aw
    n = note(s.shapes, X0, 356, CW, 128,
             "This is a classroom model, not a universal grammatical ranking. "
             "Classical Arabic rhetoric links the amount of emphasis to the "
             "listener: no doubt, no emphasis; hesitation, some emphasis; "
             "denial, stronger emphasis. Restriction (3) is a different job, "
             "not simply 'more emphasis'.", label="HOW TO READ THIS MODEL",
             size=14)
    B.click(n)


@slide("open")
def s_route(s, B, ctx):
    frame(s, "Your Route Through the Chapter", page=ctx.page)
    rows = [
        ("1", "Part 0", "What emphasis is; how it differs from restriction, "
         "exclamation and 'yes'", "key terms only", rng(ctx, 0)),
        ("2", "Part 1", "Emphasis inside a sentence", "إِنّ · أَنّ", rng(ctx, 1)),
        ("3", "Part 2", "Emphasis before and within a verb", "قَد · لَقَد · لَـ",
         rng(ctx, 2)),
        ("4", "Part 3", "Restriction, identity and totality",
         "إِنَّمَا · نَفْس · عَيْن · كُلّ · جَمِيع", rng(ctx, 3)),
        ("5", "Part 4", "Putting the tools together", "all ten tools",
         rng(ctx, 4)),
        ("6", "Part 5", "Clearing the confusion", "look-alike pairs",
         rng(ctx, 5)),
    ]
    rows = [(Para(a, size=16, color=WINE, bold=True, align="c"), b, c_, d, e)
            for a, b, c_, d, e in rows]
    table(s, B, X0, 70, [60, 90, 330, 250, 114],
          ["Stop", "Part", "Focus", "Tools Introduced", "Slides"], rows,
          rowh=56, en_size=12.5, ar_size=19)
    n = note(s.shapes, X0, 432, CW, 60, "Read every row, card and table in "
             "this chapter from the RIGHT, just like Arabic.",
             label="READING ORDER", size=14)
    B.click(n)


@slide("open")
def s_roadmap(s, B, ctx):
    frame(s, "The Chapter Road Map", page=ctx.page)
    c = s.shapes
    ry, rh = 262, 30
    centers = [826, 702, 578, 454, 330, 206]
    start_x, finish_x = 912, 128
    cards = [
        ("Stop 1 · What Emphasis Means", [("تَوْكِيد", "tawkīd — emphasis"),
                                          ("قَصْر", "qaṣr — restriction"),
                                          ("تَعَجُّب", "taʿajjub — exclamation")]),
        ("Stop 2 · Inside a Sentence", [("إِنّ", "inna — indeed"),
                                        ("أَنّ", "anna — that")]),
        ("Stop 3 · With a Verb", [("قَد", "qad — already / may"),
                                  ("لَقَد", "laqad — indeed (+ past)"),
                                  ("لَـ", "la- — truly (prefix)")]),
        ("Stop 4 · Limits & Totality", [("إِنَّمَا", "innamā — only"),
                                        ("نَفْس", "nafs — self, same"),
                                        ("عَيْن", "ʿayn — the very same"),
                                        ("كُلّ", "kull — all, every"),
                                        ("جَمِيع", "jamīʿ — all, altogether")]),
        ("Stop 5 · Putting It Together", [("", "all ten tools"),
                                          ("", "the 3-step selection method"),
                                          ("", "real-life situations")]),
        ("Stop 6 · Clearing Confusion", [("إِنّ / أَنّ", "inna / anna"),
                                         ("قَد / لَقَد", "qad / laqad"),
                                         ("نَفْس / عَيْن", "nafs / ʿayn"),
                                         ("كُلّ / جَمِيع", "kull / jamīʿ")]),
    ]
    # start flag (green) on the RIGHT
    shape(c, MSO_SHAPE.RECTANGLE, start_x + 2, ry - 46, 3, 76, "333333")
    shape(c, MSO_SHAPE.RECTANGLE, start_x + 5, ry - 46, 24, 16, FLAGGREEN)
    tb(c, start_x - 20, ry + 32, 60, 18, Para("START", size=10, color=FLAGGREEN,
       bold=True, align="c"), margins=(0, 0, 0, 0))
    pts = [start_x] + centers + [finish_x]
    w = 232
    # z-order: road first, then stop circles, then cards (builds run in
    # reading order regardless: segment, circle, card, from the right)
    segs, circs, cardg = [], [], []
    for i in range(7):
        x_right, x_left = pts[i], pts[i + 1]
        seg = shape(c, MSO_SHAPE.RECTANGLE, x_left, ry, x_right - x_left, rh, ROAD)
        dash = shape(c, MSO_SHAPE.RECTANGLE, x_left + 16, ry + rh / 2 - 1.5,
                     x_right - x_left - 32, 3, WHITE)
        segs.append((seg, dash))
    for i in range(6):
        cx = centers[i]
        circ = shape(c, MSO_SHAPE.OVAL, cx - 19, ry - 4, 38, 38, WINE, WHITE, 2)
        write_paras(circ.text_frame, [Para(str(i + 1), size=15, color=WHITE,
                    bold=True, align="c")], "m", (0, 0, 0, 0))
        circs.append(circ)
    for i in range(6):
        cx = centers[i]
        title, tools = cards[i]
        body = []
        for ar, en in tools:
            if ar:
                body.append(Para("%s   {t:%s}" % (ar, en), size=16, align="r",
                                 kind="ar", space_after=1, line=1.0))
            else:
                body.append(Para("• " + en, size=12, align="l", kind="en"))
        body.append(Para(rng(ctx, i), size=10.5, color=MUTED, align="c",
                         kind="en"))
        h = 66 + 22 * len(tools)
        bx = min(max(cx - w / 2, X0), X1 - w)
        above = i % 2 == 0
        by = ry - 14 - h if above else ry + rh + 14
        cardg.append(card(c, bx, by, w, h, title, body, title_size=12.5,
                          title_h=26, en_size=12, ar_size=16))
    for i in range(6):
        B.click(segs[i][0], "wipe")
        B.with_(segs[i][1], "wipe")
        B.click(circs[i])
        B.click(cardg[i])
    B.click(segs[6][0], "wipe")
    B.with_(segs[6][1], "wipe")
    # gold chequered finish flag with a star on the LEFT
    fg = group(c)
    shape(fg.shapes, MSO_SHAPE.RECTANGLE, finish_x - 4, ry - 50, 3, 80, "333333")
    for r_ in range(3):
        for k in range(4):
            col = GOLD if (r_ + k) % 2 == 0 else "1F1F1F"
            shape(fg.shapes, MSO_SHAPE.RECTANGLE, finish_x - 4 - 8 * (k + 1),
                  ry - 50 + 7 * r_, 8, 7, col)
    shape(fg.shapes, MSO_SHAPE.STAR_5_POINT, finish_x - 15, ry - 74, 22, 22,
          GOLD, "8A6D10", 0.75)
    tb(fg.shapes, finish_x - 46, ry + 32, 60, 18, Para("FINISH", size=10,
       color="8A6D10", bold=True, align="c"), margins=(0, 0, 0, 0))
    B.with_(fg)


@slide("open")
def s_reading_guide(s, B, ctx):
    frame(s, "How to Read the Arabic in This Chapter", page=ctx.page)
    w, gap = 272, 14
    x = X1 - w
    cards_ = [
        ("Final-Letter Convention", [
            Para("Course rule: the final letter of a word carries no vowel "
                 "sign. Case endings appear in the transliteration.",
                 size=13, space_after=8),
            Para("إِنّ الْبَيْت كَبِير", size=19, align="c", kind="ar"),
            Para("{t:inna l-bayt}{k:a}{t: kabīr}{k:un}", size=13, align="c",
                 kind="en"),
            Para("A final shadda stays, because it is part of the letter: "
                 "كُلّ · إِنّ", size=12, color=MUTED)]),
        ("Transliteration Key", [
            Para("ā ī ū  long vowels", size=13),
            Para("ʾ  hamza · ʿ  ʿayn", size=13),
            Para("ḥ ṣ ḍ ṭ ẓ  emphatic / deep sounds", size=13),
            Para("th dh kh sh gh  as in think, this, Bach, ship, Paris r",
                 size=13),
            Para("l- before a word = the article al- after a vowel: "
                 "{t:inna l-bayta}", size=12, color=MUTED)]),
        ("Colour Key", [
            Para("{b:Blue} — the emphatic tool or grammar point in focus",
                 size=13, space_after=6),
            Para("{g:✓ Green} — correct answer or correct form", size=13,
                 space_after=6),
            Para("{r:✗ Red} — incorrect form", size=13, space_after=6),
            Para("{w:Wine} — chapter headings and tool numbers", size=13,
                 space_after=6),
            Para("Labels always accompany colours.", size=12, color=MUTED)]),
    ]
    for t, body in cards_:
        g = card(s.shapes, x, 74, w, 300, t, body, title_size=15)
        B.click(g, "wipe")
        x -= w + gap
    n = note(s.shapes, X0, 388, CW, 104,
             "{k:Register labels:} FORMAL = writing, news, speeches · "
             "EVERYDAY = normal conversation in Modern Standard Arabic · "
             "CONTEXT = the meaning depends on the situation. Spoken dialects "
             "differ; where it matters, the slide says so.",
             label="REGISTER", size=14)
    B.click(n)


# ===================================================================== PART 0
def opener(s, B, ctx, n, title, focus_line, ar_title, tr_title, focus_text,
           tiles, goals, tiles_label="TOOLS IN THIS PART"):
    a, b = ctx.ranges[n - 1]
    frame(s, f"You Are Here · Stop {n} of 6", page=ctx.page,
          footer_mid=f"Stop {n} of 6 · slides {a}–{b}")
    c = s.shapes
    # six-stop navigation track, first stop at the RIGHT
    xs = [880, 756, 632, 508, 384, 260]
    shape(c, MSO_SHAPE.RECTANGLE, xs[-1], 80, xs[0] - xs[-1], 3, LINE)
    for i, x in enumerate(xs, 1):
        cur = i == n
        done = i < n
        d = 30 if cur else 22
        fill = WINE if cur else (TINT if done else WHITE)
        col = WHITE if cur else WINE
        o = shape(c, MSO_SHAPE.OVAL, x - d / 2, 81.5 - d / 2, d, d, fill,
                  WINE if (cur or done) else GREYBOX, 1.5)
        write_paras(o.text_frame, [Para(str(i), size=12 if cur else 10,
                    color=col if (cur or done) else MUTED, bold=True,
                    align="c")], "m", (0, 0, 0, 0))
        tb(c, x - 62, 98, 124, 18, Para(STOP_NAMES[i - 1], size=9.5,
           color=WINE if cur else MUTED, bold=cur, align="c"),
           margins=(0, 0, 0, 0))
    tb(c, X0, 70, 150, 22, Para("YOU ARE HERE · Stop %d" % n, size=11,
       color=WINE, bold=True), margins=(0, 0, 0, 0), anchor="m")
    # hero box (right) + focus box (left)
    hero = tb(c, 404, 124, X1 - 404, 104, [
        Para(title, size=28, color=WHITE, bold=True, align="c", kind="en",
             space_after=2),
        Para(focus_line, size=14, color=WHITE, bold=True, align="c",
             kind="en", space_after=2),
        Para("%s  {u:· %s}" % (ar_title, tr_title), size=18, color=WHITE,
             align="c", kind="ar")], fill=WINE, radius=12, anchor="m",
        margins=(10, 4, 10, 2))
    foc = tb(c, X0, 124, 300, 104, [
        Para("THE FOCUS", size=11, color=WINE, bold=True, kind="en"),
        Para(focus_text, size=12.5, kind="en")], fill=TINT, line=WINE,
        radius=12, anchor="m", margins=(12, 4, 12, 4))
    B.click(hero, "wipe")
    B.click(foc)
    tb(c, X0, 234, CW, 20, Para(tiles_label, size=14, color=WINE, bold=True,
       align="c"), anchor="m", margins=(0, 0, 0, 0))
    k = len(tiles)
    tw = 170 if k <= 4 else 160
    gap = 10
    total = k * tw + (k - 1) * gap
    x = X0 + (CW - total) / 2 + total - tw
    tl = []
    for en, ar, tr in tiles:
        tl.append(tb(c, x, 258, tw, 100, [
            Para(en, size=12, color=WINE, bold=True, align="c", kind="en"),
            Para(ar, size=20, align="c", kind="ar", space_after=0),
            Para("{t:%s}" % tr, size=11.5, align="c", kind="en")],
            fill=WHITE, line=WINE, lw=1.25, radius=10, anchor="m",
            margins=(4, 6, 4, 3)))
        x -= tw + gap
    B.chain(tl)
    gw = (CW - 2 * 12) / 3
    x = X1 - gw
    gl = []
    for i, gtxt in enumerate(goals, 1):
        gl.append(tb(c, x, 370, gw, 128, [
            Para(f"GOAL {i}", size=11, color=BLUE, bold=True, kind="en"),
            Para(gtxt, size=14, kind="en")], fill=LBLUE, line=BLUE, lw=0.75,
            radius=10, anchor="t", margins=(10, 8, 10, 4)))
        x -= gw + 12
    B.chain(gl)


@slide(0)
def p0_open(s, B, ctx):
    opener(s, B, ctx, 1, "What Emphasis Means",
           "Why speakers add force — and what counts as emphasis",
           "مَعْنَى التَّوْكِيد", "maʿnā t-tawkīd",
           "Hear the difference between a neutral and an emphasized "
           "statement, and separate emphasis from look-alike ideas.",
           [("emphasis", "تَوْكِيد", "tawkīd"),
            ("restriction", "قَصْر", "qaṣr"),
            ("exclamation", "تَعَجُّب", "taʿajjub"),
            ("emphasis of meaning", "تَوْكِيد مَعْنَوِيّ", "tawkīd maʿnawī")],
           ["Explain what emphasis adds to a sentence and why speakers use it.",
            "Tell emphasis apart from restriction, exclamation and plain 'yes'.",
            "Sort the tools by where they work: sentence, verb, attached, or word."],
           tiles_label="KEY TERMS IN THIS PART")


@slide(0)
def p0_what(s, B, ctx):
    frame(s, "What Is Emphasis (تَوْكِيد)?", stop=1, page=ctx.page)
    banner(s, B, "Emphasis ({t:tawkīd}) strengthens a statement so that the "
           "listener accepts it.")
    data = [
        ("To Remove Doubt", "{b:إِنّ} الِامْتِحَان سَهْل",
         "inna l-imtiḥāna sahlun", "The exam really is easy."),
        ("To Answer a Denial", "{b:لَقَد} دَفَعْت الْفَاتُورَة",
         "laqad dafaʿtu l-fātūrata", "I have paid the bill!"),
        ("To Stress Importance", "{b:إِنّ} سَلَامَتَكُم تَهُمُّنَا",
         "inna salāmatakum tahummunā", "Your safety truly matters to us."),
        ("To Show Feeling", "{b:لَقَد} اشْتَقْت إِلَيْك",
         "laqadi shtaqtu ilayka", "I've really missed you."),
    ]
    w, gap = 416, 12
    pos = [(X1 - w, 106), (X0, 106), (X1 - w, 294), (X0, 294)]
    for (t, ar, tr, en), (x, y) in zip(data, pos):
        g = card(s.shapes, x, y, w, 176, t, [
            Para(ar, size=20, align="c", kind="ar"),
            Para("{t:%s}" % tr, size=12.5, align="c", kind="en"),
            Para(en, size=13, align="c", kind="en", bold=True)],
            anchor="m", title_size=15)
        B.click(g, "wipe")
    legend(s, [(BLUE, LBLUE, "blue = the emphatic tool")])


@slide(0)
def p0_side(s, B, ctx):
    frame(s, "Neutral or Emphasized? Side by Side", stop=1, page=ctx.page)
    rows = [
        ("الْجَوّ جَمِيل\n~al-jawwu jamīlun\nThe weather is nice.",
         "{b:إِنّ} الْجَوّ جَمِيل\n~inna l-jawwa jamīlun\nThe weather really is nice.",
         "The whole statement is stressed."),
        ("وَصَل الضُّيُوف\n~waṣala ḍ-ḍuyūfu\nThe guests arrived.",
         "{b:لَقَد} وَصَل الضُّيُوف\n~laqad waṣala ḍ-ḍuyūfu\nThe guests have (finally) arrived.",
         "A completed event, stated with force."),
        ("أُرِيد أَن أُسَاعِد\n~urīdu an usāʿida\nI want to help.",
         "{b:إِنَّمَا} أُرِيد أَن أُسَاعِد\n~innamā urīdu an usāʿida\nI only want to help.",
         "The statement is limited (restriction)."),
        ("حَضَر الْمُدِير\n~ḥaḍara l-mudīru\nThe manager attended.",
         "حَضَر الْمُدِير {b:نَفْسُه}\n~ḥaḍara l-mudīru nafsuhu\nThe manager himself attended.",
         "The identity of the person is stressed."),
        ("نَجَح الطُّلَّاب\n~najaḥa ṭ-ṭullābu\nThe students passed.",
         "نَجَح الطُّلَّاب {b:كُلُّهُم}\n~najaḥa ṭ-ṭullābu kulluhum\nThe students all passed.",
         "Totality is stressed: no one is left out."),
    ]
    table(s, B, X0, 66, [300, 330, 214], ["Neutral", "Emphasized",
          "What Changes"], rows, rowh=76, ar_size=17, en_size=11)
    legend(s, [(BLUE, LBLUE, "blue = the emphatic tool")], y=480)


@slide(0)
def p0_five(s, B, ctx):
    frame(s, "Five Ideas Learners Often Mix Up", stop=1, page=ctx.page)
    data = [
        ("Emphasis", "تَوْكِيد · {t:tawkīd}", "Strengthens the truth of a statement.",
         "{b:إِنّ} الْعِلْم نُور", "inna l-ʿilma nūrun", "Knowledge is truly light."),
        ("Certainty of an Event", "تَحْقِيق · {t:taḥqīq}",
         "Presents an action as really done.",
         "{b:قَد} وَصَل الْقِطَار", "qad waṣala l-qiṭāru", "The train has arrived."),
        ("Confirmation", "answer words", "Agrees with what was said; a reply.",
         "نَعَم، بِالتَّأْكِيد", "naʿam, bi-t-taʾkīdi", "Yes, certainly."),
        ("Restriction", "قَصْر · {t:qaṣr}", "Limits the statement to one thing.",
         "{b:إِنَّمَا} أَنَا ضَيْف", "innamā ana ḍayfun", "I'm only a guest."),
        ("Exclamation", "تَعَجُّب · {t:taʿajjub}",
         "Expresses wonder — Chapter 2.12.",
         "مَا أَجْمَل الْبَحْر!", "mā ajmala l-baḥra!", "How beautiful the sea is!"),
    ]
    w, h, gap = 272, 196, 14
    pos = [(X1 - w, 66), (X1 - 2 * w - gap, 66), (X0, 66),
           (X1 - w - 140, 274), (X1 - 2 * w - gap - 140, 274)]
    for (t, term, what, ar, tr, en), (x, y) in zip(data, pos):
        g = card(s.shapes, x, y, w, h, t, [
            Para(term, size=13, color=WINE, align="c", kind="en"),
            Para(what, size=12, align="c", kind="en", space_after=4),
            Para(ar, size=19, align="c", kind="ar", space_after=0),
            Para("{t:%s}" % tr, size=12, align="c", kind="en"),
            Para(en, size=12, align="c", kind="en", bold=True)],
            anchor="m", title_size=15)
        B.click(g, "wipe")
    legend(s, [(BLUE, LBLUE, "blue = the emphatic tool")], y=480)


@slide(0)
def p0_places(s, B, ctx):
    frame(s, "Four Places Emphasis Can Sit", stop=1, page=ctx.page)
    banner(s, B, "Ask: does the tool work on a sentence, a verb, a single word "
           "— or is it a noun?")
    data = [
        ("On a Whole Sentence", "particles", "{b:إِنّ} · {b:أَنّ} · {b:إِنَّمَا}",
         "{b:إِنّ} الْوَقْت ثَمِين", "inna l-waqta thamīnun", "Time is precious."),
        ("On a Verb", "particles", "{b:قَد} · {b:لَقَد}",
         "{b:قَد} وَصَلْنَا", "qad waṣalnā", "We have arrived."),
        ("Attached to a Word", "a prefix", "{b:لَـ}",
         "إِنّ الْحَقّ {b:لَ}وَاضِح", "inna l-ḥaqqa la-wāḍiḥun",
         "The truth is clearly evident."),
        ("After a Noun", "nouns + pronoun suffix",
         "{b:نَفْس} · {b:عَيْن} · {b:كُلّ} · {b:جَمِيع}",
         "جَاء الضُّيُوف {b:كُلُّهُم}", "jāʾa ḍ-ḍuyūfu kulluhum",
         "The guests all came."),
    ]
    w, gap = 202, 12
    x = X1 - w
    for t, kind, tools, ar, tr, en in data:
        g = card(s.shapes, x, 104, w, 250, t, [
            Para(kind.upper(), size=10.5, color=MUTED, bold=True, align="c",
                 kind="en"),
            Para(tools, size=20, align="c", kind="ar", space_after=10),
            Para(ar, size=18, align="c", kind="ar", space_after=0),
            Para("{t:%s}" % tr, size=11.5, align="c", kind="en"),
            Para(en, size=11.5, align="c", kind="en")], anchor="m",
            title_size=14)
        B.click(g, "wipe")
        x -= w + gap
    n = note(s.shapes, X0, 366, CW, 108,
             "The words {b:نَفْس}, {b:عَيْن}, {b:كُلّ} and {b:جَمِيع} are "
             "nouns, not particles. When they follow a noun with a matching pronoun "
             "suffix, Arabic grammar calls this التَّوْكِيد الْمَعْنَوِيّ "
             "({t:at-tawkīd al-maʿnawī}, 'emphasis of meaning').",
             label="GRAMMAR NOTE", size=14)
    B.click(n)


@slide(0)
def p0_intensifiers(s, B, ctx):
    frame(s, "Grammatical Emphasis vs. Everyday Intensifiers", stop=1,
          page=ctx.page)
    banner(s, B, "Words like 'really' and 'very' intensify — but they are "
           "vocabulary, not grammatical tools.")
    rows = [
        ("حَقًّا\n~ḥaqqan", "truly, really", "هَذَا حَقًّا مُفِيد\n~hādhā ḥaqqan mufīdun\nThis is truly useful."),
        ("فِعْلًا\n~fiʿlan", "really, actually", "أَنَا فِعْلًا مُتْعَب\n~ana fiʿlan mutʿabun\nI'm really tired."),
        ("بِالتَّأْكِيد\n~bi-t-taʾkīdi", "certainly, of course", "سَأَحْضُر بِالتَّأْكِيد\n~sa-aḥḍuru bi-t-taʾkīdi\nI'll certainly come."),
        ("جِدًّا\n~jiddan", "very", "الدَّرْس مُهِمّ جِدًّا\n~ad-darsu muhimmun jiddan\nThe lesson is very important."),
    ]
    table(s, B, X0, 104, [180, 200, 464], ["Word", "Meaning", "Example"],
          rows, rowh=72, ar_size=18, en_size=11.5)
    n = note(s.shapes, X0, 436, CW, 60, "These words can combine with the "
             "chapter's tools, but they do not change grammar the way "
             "{b:إِنّ} does.", label="DIFFERENCE", size=14)
    B.click(n)


@slide(0)
def p0_notool(s, B, ctx):
    frame(s, "Emphasis Without a Tool: Repetition, Intonation, Context",
          stop=1, page=ctx.page)
    w, gap = 272, 14
    x = X1 - w
    data = [
        ("Repetition", "Repeating a word strengthens it. Grammar calls this "
         "التَّوْكِيد اللَّفْظِيّ ({t:at-tawkīd al-lafẓī}).", [
            ("نَعَم نَعَم، فَهِمْت", "naʿam naʿam, fahimtu", "Yes, yes — I understood."),
            ("مُهِمّ جِدًّا جِدًّا", "muhimmun jiddan jiddan", "Very, very important.")]),
        ("Intonation", "In speech, extra stress on one word marks it as the "
         "point. Writing cannot show this.", [
            ("الِاجْتِمَاع الْيَوْم", "al-ijtimāʿu l-YAWMA",
             "The meeting is TODAY (not tomorrow).")]),
        ("Context", "A plain sentence can sound emphatic as a reply to doubt.", [
            ("— هَل أَنْت مُتَأَكِّد؟", "hal anta mutaʾakkidun?", "Are you sure?"),
            ("— الْمَتْحَف مَفْتُوح.", "al-matḥafu maftūḥun.",
             "The museum IS open.")]),
    ]
    for t, expl, exs in data:
        body = [Para(expl, size=12.5, kind="en", space_after=8)]
        for ar, tr, en in exs:
            body += [Para(ar, size=18, align="r", kind="ar", space_after=0),
                     Para("{t:%s}  —  %s" % (tr, en), size=11.5, align="r",
                          kind="en", space_after=8)]
        g = card(s.shapes, x, 70, w, 330, t, body, title_size=15, en_size=13,
                 ar_size=19, anchor="m")
        B.click(g, "wipe")
        x -= w + gap
    n = note(s.shapes, X0, 412, CW, 80, "Capital letters in the "
             "transliteration show spoken stress only. The Arabic spelling "
             "does not change.", label="NOTE", size=14)
    B.click(n)


@slide(0)
def p0_context(s, B, ctx):
    frame(s, "Why Context Matters", stop=1, page=ctx.page)
    banner(s, B, "The same tool can do different jobs. Read the whole situation.")
    data = [
        ("Same Particle, Different Verb", [
            ("{b:قَد} وَصَل الْقِطَار", "qad waṣala l-qiṭāru",
             "The train has arrived. (completed)"),
            ("{b:قَد} يَصِل الْقِطَار", "qad yaṣilu l-qiṭāru",
             "The train may arrive. (possible)")]),
        ("Same Sentence, Different Listener", [
            ("الْفُنْدُق قَرِيب", "al-funduqu qarībun",
             "The hotel is near. (just information)"),
            ("{b:إِنّ} الْفُنْدُق قَرِيب", "inna l-funduqa qarībun",
             "The hotel IS near. (to someone who doubts)")]),
        ("Same Word, Different Job", [
            ("{b:نَفْس} الشَّخْص", "nafsu sh-shakhṣi", "the same person"),
            ("الشَّخْص {b:نَفْسُه}", "ash-shakhṣu nafsuhu",
             "the person himself")]),
    ]
    w, gap = 272, 14
    x = X1 - w
    for t, exs in data:
        body = []
        for ar, tr, en in exs:
            body += [Para(ar, size=19, align="c", kind="ar", space_after=0),
                     Para("{t:%s}" % tr, size=12, align="c", kind="en",
                          space_after=0),
                     Para(en, size=12, align="c", kind="en", space_after=12)]
        g = card(s.shapes, x, 106, w, 300, t, body, title_size=14.5,
                 anchor="m")
        B.click(g, "wipe")
        x -= w + gap
    legend(s, [(BLUE, LBLUE, "blue = the tool whose job changes")], y=420)


@slide(0)
def p0_check(s, B, ctx):
    frame(s, "Quick Check: Spot the Emphasis", stop=1, page=ctx.page)
    drill_header(s, y=66, labels=("Sentence", "Look at…", "What kind?"))
    items = [
        dict(q=ex("إِنّ الصِّحَّة كَنْز", "inna ṣ-ṣiḥḥata kanzun",
                  "Health is truly a treasure."),
             cue=Para("إِنّ", size=19, kind="ar", align="c", color=BLUE),
             a="{g:✓ Emphasis}\nof the whole statement"),
        dict(q=ex("حَضَر الْوَزِير نَفْسُه", "ḥaḍara l-wazīru nafsuhu",
                  "The minister himself attended."),
             cue=Para("نَفْسُه", size=19, kind="ar", align="c", color=BLUE),
             a="{g:✓ Identity}\nemphasis on the person"),
        dict(q=ex("مَا أَجْمَل هَذَا الْمَكَان!", "mā ajmala hādhā l-makāna!",
                  "How beautiful this place is!"),
             cue=Para("مَا أَجْمَل", size=19, kind="ar", align="c", color=BLUE),
             a="{g:✓ Exclamation}\nnot emphasis (Ch. 2.12)"),
        dict(q=ex("إِنَّمَا أَنَا طَالِب", "innamā ana ṭālibun",
                  "I'm only a student."),
             cue=Para("إِنَّمَا", size=19, kind="ar", align="c", color=BLUE),
             a="{g:✓ Restriction}\nlimits the statement"),
        dict(q=ex("قَرَأْت الرِّسَالَة كُلَّهَا", "qaraʾtu r-risālata kullahā",
                  "I read the whole letter."),
             cue=Para("كُلَّهَا", size=19, kind="ar", align="c", color=BLUE),
             a="{g:✓ Totality}\nnothing is left out"),
    ]
    drill(s, B, items, y0=90, rowh=70, gap=8)


# ===================================================================== PART 1
@slide(1)
def p1_open(s, B, ctx):
    opener(s, B, ctx, 2, "Emphasis Inside a Sentence",
           "Strengthening a whole statement — and reporting one",
           "التَّوْكِيد فِي الْجُمْلَة الِاسْمِيَّة", "at-tawkīd fī l-jumla l-ismiyya",
           "Open a nominal sentence with إِنّ, apply its effect on the "
           "endings, and use أَنّ for a fact inside a larger sentence.",
           [("indeed, truly", "إِنّ", "inna"), ("that", "أَنّ", "anna")],
           ["Use إِنّ to open an emphatic nominal sentence and apply its "
            "case effect.",
            "Attach pronouns correctly: إِنَّه، إِنَّهَا، إِنَّنِي.",
            "Choose between إِنّ and أَنّ by their position in the sentence."])


@slide(1)
def p1_inna(s, B, ctx):
    frame(s, "Tool 1 · إِنّ (inna) — Indeed, Truly", stop=2, page=ctx.page)
    w = 410
    facts = card(s.shapes, X1 - w, 70, w, 300, "What It Is", [
        Para("{k:Type:} a particle ({t:ḥarf}) — one of the 'sisters of "
             "{t:inna}'", size=13, space_after=6),
        Para("{k:Position:} at the start of a nominal sentence (a sentence "
             "that begins with a noun)", size=13, space_after=6),
        Para("{k:Meaning:} presents the whole statement as certain", size=13,
             space_after=6),
        Para("{k:Grammar:} the subject takes the accusative ending {t:-a}; "
             "the predicate keeps its nominative ending", size=13,
             space_after=6),
        Para("{k:Register:} FORMAL and written Arabic; set phrases in speech",
             size=13)], title_size=15)
    B.click(facts, "wipe")
    y = 70
    for ar, tr, en in [
            ("{b:إِنّ} الْعِلْم نُور", "inna l-ʿilma nūrun", "Knowledge is truly light."),
            ("{b:إِنّ} اللُّغَة الْعَرَبِيَّة جَمِيلَة",
             "inna l-lughata l-ʿarabiyyata jamīlatun",
             "The Arabic language is truly beautiful."),
            ("{b:إِنّ} الِاجْتِمَاع مُهِمّ", "inna l-ijtimāʿa muhimmun",
             "The meeting is indeed important.")]:
        b = exbox(s.shapes, X0, y, CW - w - 14, 92, ar, tr, en)
        B.click(b, "wipe")
        y += 104
    n = note(s.shapes, X0, 386, CW, 92, "Do not confuse it with إِنْ "
             "({t:in}, 'if'), as in إِن شَاء اللَّه ({t:in shāʾa llāhu}, "
             "'if God wills'). Tool 1 is always {t:inna} with a doubled n.",
             label="LOOK-ALIKE", size=14)
    B.click(n)


@slide(1)
def p1_refresher(s, B, ctx):
    frame(s, "Nominal Sentence Refresher: Subject + Predicate", stop=2,
          page=ctx.page)
    banner(s, B, "A nominal sentence = a subject (مُبْتَدَأ) + a predicate (خَبَر).")
    c = s.shapes
    # without inna
    row1 = group(c)
    tb(row1.shapes, 620, 112, 260, 86, [Para("الْبَيْت", size=24, align="c",
       kind="ar"), Para("{t:al-bayt}{k:u} · subject (nominative)", size=12,
       align="c", kind="en")], fill=WHITE, line=WINE, radius=10, anchor="m")
    tb(row1.shapes, 330, 112, 260, 86, [Para("كَبِير", size=24, align="c",
       kind="ar"), Para("{t:kabīr}{k:un} · predicate (nominative)", size=12,
       align="c", kind="en")], fill=WHITE, line=WINE, radius=10, anchor="m")
    tb(row1.shapes, X0, 112, 220, 86, Para("The house is big.", size=14,
       align="c"), anchor="m")
    B.click(row1, "wipe")
    lab = tb(c, X0, 210, CW, 30, Para("Add إِنّ at the start ↓", size=15,
             color=BLUE, bold=True, align="c"), anchor="m")
    B.click(lab)
    row2 = group(c)
    tb(row2.shapes, 820, 252, 120, 96, Para("{b:إِنّ}", size=26, align="c",
       kind="ar"), fill=LBLUE, line=BLUE, radius=10, anchor="m")
    tb(row2.shapes, 540, 252, 260, 96, [Para("الْبَيْت", size=24, align="c",
       kind="ar"), Para("{t:al-bayt}{b:a} · 'subject of inna' (accusative)",
       size=12, align="c", kind="en")], fill=WHITE, line=BLUE, radius=10,
       anchor="m")
    tb(row2.shapes, 270, 252, 260, 96, [Para("كَبِير", size=24, align="c",
       kind="ar"), Para("{t:kabīr}{k:un} · 'predicate of inna' (nominative)",
       size=12, align="c", kind="en")], fill=WHITE, line=WINE, radius=10,
       anchor="m")
    tb(row2.shapes, X0, 252, 166, 96, Para("The house is indeed big.",
       size=14, align="c"), anchor="m")
    B.click(row2, "wipe")
    n = note(c, X0, 368, CW, 110,
             "Only the subject changes: {t:-u} becomes {t:-a}. The predicate keeps "
             "{t:-un}. Arabic grammar names them اِسْم إِنّ ({t:ism inna}) and "
             "خَبَر إِنّ ({t:khabar inna}).", label="THE RULE", size=15)
    B.click(n)


@slide(1)
def p1_cases(s, B, ctx):
    frame(s, "What إِنّ Does to the Case Endings", stop=2, page=ctx.page)
    rows = [
        ("الْبَيْت كَبِير\n~al-baytu kabīrun",
         "إِنّ الْبَيْت كَبِير\n~inna l-bayta kabīrun",
         "-u becomes {b:-a}\n(seen in transliteration)", "-un stays"),
        ("الْمُهَنْدِسُون مَشْغُولُون\n~al-muhandisūna mashghūlūna",
         "إِنّ الْمُهَنْدِسِ{b:ي}ن مَشْغُولُون\n~inna l-muhandisīna mashghūlūna",
         "-ūna becomes {b:-īna}\n(visible in the letters)", "-ūna stays"),
        ("الْمُدِيرَان مُسَافِرَان\n~al-mudīrāni musāfirāni",
         "إِنّ الْمُدِيرَ{b:يْ}ن مُسَافِرَان\n~inna l-mudīrayni musāfirāni",
         "-āni becomes {b:-ayni}\n(visible in the letters)", "-āni stays"),
    ]
    table(s, B, X0, 66, [250, 290, 180, 124], ["Without إِنّ", "With إِنّ",
          "Subject", "Predicate"], rows, rowh=84, ar_size=18, en_size=11.5)
    tb(s.shapes, X0, 348, CW, 24, Para("Translations: The house is big. · "
       "The engineers are busy. · The two managers are travelling.", size=12,
       color=MUTED, align="c"), anchor="m")
    n = note(s.shapes, X0, 378, CW, 96,
             "Because this course writes no vowel on the final letter, check "
             "the transliteration for {t:-a}. With sound masculine plurals "
             "and duals, the spelling itself changes: ـُون becomes ـِين, and "
             "ـَان becomes ـَيْن.", label="SPOT THE CHANGE", size=14)
    B.click(n)
    legend(s, [(BLUE, LBLUE, "blue = the ending changed by إِنّ")], y=480)


@slide(1)
def p1_pronouns(s, B, ctx):
    frame(s, "Attaching Pronouns to إِنّ", stop=2, page=ctx.page)
    rows = [
        ("I", "إِنَّنِي · إِنِّي", "innanī · innī", "إِنَّنِي سَعِيد بِلِقَائِك\n~innanī saʿīdun bi-liqāʾika\nI'm truly happy to meet you."),
        ("you (m.)", "إِنَّك", "innaka", "إِنَّك عَلَى حَقّ\n~innaka ʿalā ḥaqqin\nYou're right."),
        ("he / it", "إِنَّه", "innahu", "إِنَّه طَالِب مُجْتَهِد\n~innahu ṭālibun mujtahidun\nHe's a truly hard-working student."),
        ("she / it", "إِنَّهَا", "innahā", "إِنَّهَا مُعَلِّمَة مُمْتَازَة\n~innahā muʿallimatun mumtāzatun\nShe's truly an excellent teacher."),
        ("we", "إِنَّنَا · إِنَّا", "innanā · innā", "إِنَّنَا جَاهِزُون\n~innanā jāhizūna\nWe're ready."),
        ("they (m.)", "إِنَّهُم", "innahum", "إِنَّهُم فِي الطَّرِيق\n~innahum fī ṭ-ṭarīqi\nThey're on their way."),
    ]
    def two(d):
        ar, tr, en = d.split("\n")
        return [ar, Para("{t:%s}  —  %s" % (tr[1:], en), size=11.5,
                         align="c", kind="en")]
    rows = [(a, b, "{t:%s}" % c_, two(d)) for a, b, c_, d in rows]
    table(s, B, X0, 66, [100, 150, 130, 464], ["Pronoun", "Form",
          "Sound", "Example"], rows, rowh=56, ar_size=18, en_size=11.5)
    n = note(s.shapes, X0, 438, CW, 56, "After إِنّ the pronoun must be "
             "attached: إِنَّه مَشْغُول ✓ — never إِنّ هُوَ مَشْغُول ✗. "
             "'You (f.)' is إِنَّك too: {t:innaki}.", label="RULE", size=13)
    B.click(n)


@slide(1)
def p1_patterns(s, B, ctx):
    frame(s, "Common Patterns with إِنّ", stop=2, page=ctx.page)
    data = [
        ("Pattern 1 · إِنّ + Noun + Description", "{b:إِنّ} الْوَقْت ثَمِين",
         "inna l-waqta thamīnun", "Time is precious.",
         "The most common pattern."),
        ("Pattern 2 · إِنّ + Pronoun + Verb", "{b:إِنَّه} يَعْمَل كَثِيرًا",
         "innahu yaʿmalu kathīran", "He really works a lot.",
         "The predicate can be a whole verbal sentence."),
        ("Pattern 3 · إِنّ + Place Phrase + Noun", "{b:إِنّ} فِي الْمَكْتَب ضُيُوفًا",
         "inna fī l-maktabi ḍuyūfan", "There are guests in the office.",
         "The subject comes late but is still accusative: {t:ḍuyūf}{b:an}."),
        ("Pattern 4 · Reported Speech: قَال + إِنّ",
         "قَال الطَّبِيب {b:إِنّ} الْحَالَة مُسْتَقِرَّة",
         "qāla ṭ-ṭabību inna l-ḥālata mustaqirratun",
         "The doctor said that the condition is stable.",
         "After قَال 'said', standard grammar uses إِنّ."),
    ]
    w, h = 416, 196
    pos = [(X1 - w, 66), (X0, 66), (X1 - w, 274), (X0, 274)]
    for (t, ar, tr, en, n_), (x, y) in zip(data, pos):
        g = card(s.shapes, x, y, w, h, t, [
            Para(ar, size=19, align="c", kind="ar", space_after=0),
            Para("{t:%s}" % tr, size=12, align="c", kind="en"),
            Para(en, size=12.5, align="c", kind="en", bold=True, space_after=6),
            Para(n_, size=12, align="c", kind="en", color=MUTED)],
            anchor="m", title_size=15)
        B.click(g, "wipe")
    legend(s, [(BLUE, LBLUE, "blue = إِنّ and the ending it causes")], y=480)


@slide(1)
def p1_anna(s, B, ctx):
    frame(s, "Tool 2 · أَنّ (anna) — That", stop=2, page=ctx.page)
    w = 410
    facts = card(s.shapes, X1 - w, 70, w, 300, "What It Is", [
        Para("{k:Type:} a particle — also a 'sister of {t:inna}'", size=13,
             space_after=6),
        Para("{k:Position:} never first; it opens a clause inside a larger "
             "sentence", size=13, space_after=6),
        Para("{k:Meaning:} usually 'that'. Its emphasis is light; English "
             "simply says 'that'.", size=13, space_after=6),
        Para("{k:Grammar:} same effect as إِنّ — subject {t:-a}, predicate "
             "nominative", size=13, space_after=6),
        Para("{k:Relationship:} same family as إِنّ, but the position "
             "decides which one you use", size=13)], title_size=15)
    B.click(facts, "wipe")
    y = 70
    for ar, tr, en in [
            ("عَرَفْت {b:أَنّ} الْمَتْحَف مُغْلَق", "ʿaraftu anna l-matḥafa mughlaqun",
             "I found out that the museum is closed."),
            ("أَظُنّ {b:أَنّ} الْقِطَار مُتَأَخِّر", "aẓunnu anna l-qiṭāra mutaʾakhkhirun",
             "I think (that) the train is late."),
            ("يَسُرُّنِي {b:أَنَّكُم} هُنَا", "yasurrunī annakum hunā",
             "I'm pleased that you are here.")]:
        b = exbox(s.shapes, X0, y, CW - w - 14, 92, ar, tr, en)
        B.click(b, "wipe")
        y += 104
    n = note(s.shapes, X0, 386, CW, 92, "Before a verb, 'to / that' is أَنْ "
             "({t:an}) — a different particle: أُرِيد أَن أُسَافِر "
             "({t:urīdu an usāfira}, 'I want to travel').", label="LOOK-ALIKE",
             size=14)
    B.click(n)


@slide(1)
def p1_anna_where(s, B, ctx):
    frame(s, "Where أَنّ Appears Inside a Larger Sentence", stop=2,
          page=ctx.page)
    data = [
        ("After Verbs of Knowing, Thinking, Hearing",
         "عَرَف · عَلِم · ظَنّ · سَمِع · لَاحَظ",
         "ʿarafa · ʿalima · ẓanna · samiʿa · lāḥaẓa",
         "سَمِعْت {b:أَنّ} الطَّقْس سَيَكُون حَارًّا",
         "samiʿtu anna ṭ-ṭaqsa sa-yakūnu ḥārran",
         "I heard that the weather will be hot."),
        ("After Prepositions", "لِأَنّ · بِمَا أَنّ · رَغْم أَنّ",
         "li-anna · bi-mā anna · raghma anna",
         "لَم أَحْضُر {b:لِأَنَّنِي} كُنْت مَرِيضًا",
         "lam aḥḍur li-annanī kuntu marīḍan",
         "I didn't attend because I was ill."),
        ("After Set Phrases", "مِن الْوَاضِح أَنّ · مِن الْمُهِمّ أَنّ",
         "mina l-wāḍiḥi anna · mina l-muhimmi anna",
         "مِن الْوَاضِح {b:أَنّ} الْخُطَّة نَاجِحَة",
         "mina l-wāḍiḥi anna l-khuṭṭata nājiḥatun",
         "It is clear that the plan is working."),
    ]
    y = 68
    for t, words, wtr, ar, tr, en in data:
        g = group(s.shapes)
        tb(g.shapes, X1 - 330, y, 330, 132, [
            Para(t, size=14, color=WINE, bold=True, align="r", kind="en"),
            Para(words, size=18, align="r", kind="ar", space_after=0),
            Para("{t:%s}" % wtr, size=11.5, align="r", kind="en")],
            fill=TINT, line=WINE, radius=10, anchor="m")
        exbox(g.shapes, X0, y, CW - 342, 132, ar, tr, en)
        B.click(g, "wipe")
        y += 142
    legend(s, [(BLUE, LBLUE, "blue = أَنّ (with any attached pronoun)")],
           y=482)


@slide(1)
def p1_choose(s, B, ctx):
    frame(s, "Choosing Between إِنّ and أَنّ: Position Decides", stop=2,
          page=ctx.page)
    banner(s, B, "They are not interchangeable. Look at what comes BEFORE the "
           "particle.")
    w = 416
    a = card(s.shapes, X1 - w, 106, w, 262, "Use inna · إِنّ", [
        Para("• at the very start of a sentence", size=14),
        Para("{b:إِنّ} الصَّبْر جَمِيل", size=19, align="r", kind="ar", space_after=0),
        Para("{t:inna ṣ-ṣabra jamīlun} — Patience is beautiful.", size=12,
             align="r", space_after=8),
        Para("• after قَال 'he said' (and its forms)", size=14),
        Para("قَالَت {b:إِنّ} الطَّعَام جَاهِز", size=19, align="r", kind="ar",
             space_after=0),
        Para("{t:qālat inna ṭ-ṭaʿāma jāhizun} — She said (that) the food is "
             "ready.", size=12, align="r")], title_size=16)
    b = card(s.shapes, X0, 106, w, 262, "Use anna · أَنّ", [
        Para("• after verbs of knowing, thinking, hearing", size=14),
        Para("أَعْرِف {b:أَنّ} الطَّعَام جَاهِز", size=19, align="r", kind="ar",
             space_after=0),
        Para("{t:aʿrifu anna ṭ-ṭaʿāma jāhizun} — I know (that) the food is "
             "ready.", size=12, align="r", space_after=8),
        Para("• after a preposition or set phrase", size=14),
        Para("{b:لِأَنّ} الطَّعَام جَاهِز", size=19, align="r", kind="ar",
             space_after=0),
        Para("{t:li-anna ṭ-ṭaʿāma jāhizun} — because the food is ready",
             size=12, align="r")], title_size=16, color=BLUE, body_fill=LBLUE)
    B.click(a, "wipe")
    B.click(b, "wipe")
    n = note(s.shapes, X0, 380, CW, 96, "Modern newspapers sometimes write "
             "قَال أَنّ; standard grammar and exams expect قَال إِنّ. "
             "Learn the standard form.", label="FORMAL vs. MODERN USAGE",
             size=14)
    B.click(n)


@slide(1)
def p1_register(s, B, ctx):
    frame(s, "Formal and Everyday Use", stop=2, page=ctx.page)
    w = 416
    a = card(s.shapes, X1 - w, 70, w, 300, "FORMAL: Writing, News, Speeches", [
        Para("In formal Arabic, إِنّ often opens a statement.", size=13,
             space_after=8),
        Para("{b:إِنّ} هَذَا الْقَرَار خُطْوَة مُهِمَّة", size=19, align="r",
             kind="ar", space_after=0),
        Para("{t:inna hādhā l-qarāra khuṭwatun muhimmatun}", size=12,
             align="r"),
        Para("This decision is an important step.", size=12, align="r",
             space_after=10),
        Para("{b:إِنّ} الْحَمْد لِلَّه", size=19, align="r", kind="ar",
             space_after=0),
        Para("{t:inna l-ḥamda li-llāhi} — Praise be to God. (a fixed opening "
             "in speeches)", size=12, align="r")], title_size=15)
    b = card(s.shapes, X0, 70, w, 300, "EVERYDAY: Conversation", [
        Para("Starting a spoken sentence with إِنّ is rare. Speakers "
             "use stress or words such as:", size=13, space_after=8),
        Para("أَكِيد", size=19, align="r", kind="ar", space_after=0),
        Para("{t:akīd} — for sure", size=12, align="r", space_after=6),
        Para("وَاللَّه", size=19, align="r", kind="ar", space_after=0),
        Para("{t:wallāhi} — honestly (literally 'by God')", size=12,
             align="r", space_after=6),
        Para("In many dialects a form of إِنّ simply means 'that': "
             "Egyptian {t:inn}, Levantine {t:innu}.", size=12, color=MUTED)],
        title_size=15, color=BLUE, body_fill=LBLUE)
    B.click(a, "wipe")
    B.click(b, "wipe")
    n = note(s.shapes, X0, 384, CW, 92, "The phrase إِن شَاء اللَّه ({t:in "
             "shāʾa llāhu}, 'God willing') uses إِنْ 'if', not Tool 1. Similar "
             "letters, different word.", label="EVERYDAY LOOK-ALIKE", size=14)
    B.click(n)


@slide(1)
def p1_errors(s, B, ctx):
    frame(s, "Common Learner Errors with إِنّ and أَنّ", stop=2, page=ctx.page)
    rows = [
        ("إِنّ الْجَوّ جَمِيلًا", ["إِنّ الْجَوّ جَمِيل", "~inna l-jawwa jamīlun"],
         "The predicate stays nominative: {t:jamīlun}, not {t:jamīlan}."),
        ("إِنّ هُوَ مَشْغُول", ["إِنَّه مَشْغُول", "~innahu mashghūlun"],
         "Attach the pronoun to إِنّ."),
        ("عَرَفْت إِنّ الدَّرْس سَهْل", ["عَرَفْت أَنّ الدَّرْس سَهْل",
         "~ʿaraftu anna d-darsa sahlun"], "After a verb of knowing: أَنّ."),
        ("أُرِيد أَنّ أُسَافِر", ["أُرِيد أَن أُسَافِر", "~urīdu an usāfira"],
         "Before a verb: أَنْ ({t:an}), not أَنّ."),
        ("إِنّ الْمُعَلِّمُون مُتَعَاوِنُون", ["إِنّ الْمُعَلِّمِين مُتَعَاوِنُون",
         "~inna l-muʿallimīna mutaʿāwinūna"], "Plural subject: ـُون becomes ـِين."),
    ]
    y = 66
    for wrong, right, why in rows:
        fixrow(s, B, y, wrong, right, why, h=78)
        y += 84


@slide(1)
def p1_dialogue(s, B, ctx):
    frame(s, "Dialogue: At the Office", stop=2, page=ctx.page)
    tb(s.shapes, X0, 62, CW, 24, Para("EVERYDAY · Sāra and ʿUmar talk about a "
       "client. Spot the blue tools.", size=12.5, color=MUTED, align="r"),
       margins=(0, 0, 0, 0))
    lines = [
        ("r", "سَارَة · Sāra", "هَل صَحِيح {b:أَنّ} الْعَمِيل غَاضِب؟",
         "hal ṣaḥīḥun anna l-ʿamīla ghāḍibun?",
         "Is it true that the client is angry?"),
        ("l", "عُمَر · ʿUmar", "لَا، {b:إِنّ} الْعَمِيل سَعِيد بِالنَّتَائِج.",
         "lā, inna l-ʿamīla saʿīdun bi-n-natāʾiji.",
         "No — the client is actually happy with the results."),
        ("r", "سَارَة · Sāra", "وَلَكِن الْمُدِير قَال {b:إِنّ} هُنَاك مُشْكِلَة.",
         "wa-lākini l-mudīru qāla inna hunāka mushkilatan.",
         "But the manager said (that) there is a problem."),
        ("l", "عُمَر · ʿUmar", "نَعَم، مُشْكِلَة صَغِيرَة، وَ{b:قَد} حَلَلْنَاهَا.",
         "naʿam, mushkilatun ṣaghīratun, wa-qad ḥalalnāhā.",
         "Yes, a small problem — and we have already solved it."),
    ]
    y = 92
    for side, name, ar, tr, en in lines:
        bubble(s, B, y, side, name, ar, tr, en, h=90, w=600)
        y += 98
    legend(s, [(BLUE, LBLUE, "blue = emphatic tool (قَد previews Part 2)")],
           y=484)


@slide(1)
def p1_practice(s, B, ctx):
    frame(s, "Guided Practice: إِنّ or أَنّ?", stop=2, page=ctx.page)
    drill_header(s, y=62, labels=("Sentence", "Position cue", "Answer"))
    B_ = "ــــــ"
    items = [
        dict(q=[f"{B_} الصَّبْر مِفْتَاح الْفَرَج", "Patience is the key to relief. (proverb)"],
             cue="start of the sentence",
             a=["{g:إِنّ}", "~inna ṣ-ṣabra miftāḥu l-faraji"]),
        dict(q=[f"عَلِمْت {B_} الرِّحْلَة مُلْغَاة", "I learned that the flight is cancelled."],
             cue="after a verb of knowing",
             a=["{g:أَنّ}", "~ʿalimtu anna r-riḥlata mulghātun"]),
        dict(q=[f"قَالَت الْمُعَلِّمَة {B_} الِامْتِحَان سَهْل", "The teacher said the exam is easy."],
             cue="after قَال 'said'",
             a=["{g:إِنّ}", "~qālati l-muʿallimatu inna l-imtiḥāna sahlun"]),
        dict(q=[f"تَأَخَّرْت {B_} الطَّرِيق مُزْدَحِم", "I was late because the road is busy."],
             cue="'because' = preposition + ?",
             a=["{g:لِأَنّ}", "~taʾakhkhartu li-anna ṭ-ṭarīqa muzdaḥimun"]),
        dict(q=[f"{B_} مُوَظَّف أَمِين", "He is truly an honest employee."],
             cue="start + 'he' attached",
             a=["{g:إِنَّه}", "~innahu muwaẓẓafun amīnun"]),
        dict(q=[f"أَظُنّ {B_} الْمَطَر سَيَتَوَقَّف قَرِيبًا", "I think the rain will stop soon."],
             cue="after a verb of thinking",
             a=["{g:أَنّ}", "~aẓunnu anna l-maṭara sa-yatawaqqafu qarīban"]),
    ]
    drill(s, B, items, y0=86, rowh=64, gap=5)


@slide(1)
def p1_cumulative(s, B, ctx):
    frame(s, "Cumulative Drill: Rewrite with إِنّ or أَنّ", stop=2,
          page=ctx.page)
    drill_header(s, y=62, labels=("Rewrite this", "Instruction", "Answer"),
                 widths=(30, 300, 200, 294))
    items = [
        dict(q=["الطَّعَام لَذِيذ", "The food is delicious."],
             cue="add إِنّ", a=["{g:إِنّ الطَّعَام لَذِيذ}",
                              "~inna ṭ-ṭaʿāma ladhīdhun"]),
        dict(q=["الْمُوَظَّفُون مُجْتَهِدُون", "The employees are hard-working."],
             cue="add إِنّ — watch the plural",
             a=["{g:إِنّ الْمُوَظَّفِين مُجْتَهِدُون}",
                "~inna l-muwaẓẓafīna mujtahidūna"]),
        dict(q=["هِيَ مُتْعَبَة", "She is tired."],
             cue="add إِنّ — attach the pronoun",
             a=["{g:إِنَّهَا مُتْعَبَة}", "~innahā mutʿabatun"]),
        dict(q=["سَمِعْت: الْمَتْجَر مَفْتُوح", "I heard: the shop is open."],
             cue="join with أَنّ",
             a=["{g:سَمِعْت أَنّ الْمَتْجَر مَفْتُوح}",
                "~samiʿtu anna l-matjara maftūḥun"]),
        dict(q=["قَال الْمُدِير: الْعَرْض نَاجِح", "The manager said: the offer is successful."],
             cue="join — after قَال",
             a=["{g:قَال الْمُدِير إِنّ الْعَرْض نَاجِح}",
                "~qāla l-mudīru inna l-ʿarḍa nājiḥun"]),
    ]
    drill(s, B, items, y0=86, rowh=76, gap=6, widths=(30, 300, 200, 294))


# ===================================================================== PART 2
@slide(2)
def p2_open(s, B, ctx):
    opener(s, B, ctx, 3, "Emphasis Before and Within a Verb",
           "Completed events, possible events, and the emphatic prefix",
           "التَّوْكِيد مَعَ الْفِعْل", "at-tawkīd maʿa l-fiʿl",
           "Let the verb form tell you what قَد means, use لَقَد where it "
           "fits, and attach the prefix لَـ in the right place.",
           [("already / may", "قَد", "qad"),
            ("indeed (+ past verb)", "لَقَد", "laqad"),
            ("truly (a prefix)", "لَـ", "la-")],
           ["Read قَد correctly: past verb = completion; present verb = "
            "possibility.",
            "Use لَقَد with past verbs only — and only when emphasis fits.",
            "Attach لَـ correctly, including the pattern إِنّ … لَـ."])


@slide(2)
def p2_qad(s, B, ctx):
    frame(s, "Tool 3 · قَد (qad) — Two Meanings", stop=3, page=ctx.page)
    banner(s, B, "The particle قَد sits directly before a verb. The verb form "
           "decides its meaning.")
    w, aw = 400, 44
    a = card(s.shapes, X1 - w, 106, w, 262, "With a PAST Verb: Already, Indeed", [
        Para("Completion and certainty ({t:taḥqīq})", size=13, align="c",
             color=MUTED, space_after=8),
        Para("{b:قَد} وَصَل الْقِطَار", size=19, align="r", kind="ar", space_after=0),
        Para("{t:qad waṣala l-qiṭāru} — The train has arrived.", size=12,
             align="r", space_after=10),
        Para("{b:قَد} أَرْسَلْت التَّقْرِير", size=19, align="r", kind="ar",
             space_after=0),
        Para("{t:qad arsaltu t-taqrīra} — I have (already) sent the report.",
             size=12, align="r")], title_size=14.5)
    b = card(s.shapes, X0, 106, w, 262, "With a PRESENT Verb: May, Sometimes", [
        Para("Possibility or occasional action", size=13, align="c",
             color=MUTED, space_after=8),
        Para("{b:قَد} يَتَأَخَّر الْقِطَار", size=19, align="r", kind="ar",
             space_after=0),
        Para("{t:qad yataʾakhkharu l-qiṭāru} — The train may be late.",
             size=12, align="r", space_after=10),
        Para("{b:قَد} أَعْمَل مِن الْبَيْت غَدًا", size=19, align="r",
             kind="ar", space_after=0),
        Para("{t:qad aʿmalu mina l-bayti ghadan} — I may work from home "
             "tomorrow.", size=12, align="r")], title_size=14.5, color=BLUE,
        body_fill=LBLUE)
    B.click(a, "wipe")
    B.click(b, "wipe")
    n = note(s.shapes, X0, 382, CW, 96, "In classical and Qurʾanic texts, قَد "
             "+ present can also confirm a fact. In modern everyday and "
             "media Arabic, read it as 'may'. CONTEXT decides.",
             label="CONTEXT", size=14)
    B.click(n)


@slide(2)
def p2_verbdecides(s, B, ctx):
    frame(s, "The Verb Decides: قَد + Past vs. قَد + Present", stop=3,
          page=ctx.page)
    rows = [
        ("قَد + past verb", "قَد فَهِمْت\n~qad fahimtu", "completion, certainty",
         "I have understood."),
        ("قَد + present verb", "قَد أَفْهَم لَاحِقًا\n~qad afhamu lāḥiqan",
         "possibility", "I may understand later."),
        ("قَد + present verb", "قَد يَغْضَب أَحْيَانًا\n~qad yaghḍabu aḥyānan",
         "occasional action", "He can sometimes get angry."),
        ("قَد لَا + present verb", "قَد لَا يَأْتِي\n~qad lā yaʾtī",
         "negative possibility", "He may not come."),
    ]
    table(s, B, X0, 66, [190, 250, 190, 214], ["Form", "Example",
          "Meaning", "Translation"], rows, rowh=64, ar_size=18, en_size=12)
    w = 416
    a = note(s.shapes, X1 - w, 368, w, 116, "The particle قَد sits directly "
             "before the verb. Only لَا 'not' may come between: قَد لَا.",
             label="PLACEMENT", size=14)
    b = note(s.shapes, X0, 368, w, 116, "Future سَـ or سَوْفَ never follows "
             "قَد: {r:✗} قَد سَيَذْهَب. Say سَيَذْهَب ({t:sa-yadhhabu}) or "
             "قَد يَذْهَب ({t:qad yadhhabu}).", label="NOT USED", size=14,
             color=RED, fill=LRED)
    B.click(a)
    B.click(b)


@slide(2)
def p2_laqad(s, B, ctx):
    frame(s, "Tool 4 · لَقَد (laqad) — Indeed (with a Past Verb)", stop=3,
          page=ctx.page)
    c = s.shapes
    g = group(c)
    tb(g.shapes, 700, 70, 120, 70, Para("{b:لَـ}", size=26, align="c",
       kind="ar"), fill=LBLUE, line=BLUE, radius=10, anchor="m")
    tb(g.shapes, 660, 86, 40, 40, Para("+", size=22, align="c"), anchor="m")
    tb(g.shapes, 540, 70, 120, 70, Para("{b:قَد}", size=26, align="c",
       kind="ar"), fill=LBLUE, line=BLUE, radius=10, anchor="m")
    tb(g.shapes, 500, 86, 40, 40, Para("=", size=22, align="c"), anchor="m")
    tb(g.shapes, 360, 70, 140, 70, Para("{w:لَقَد}", size=28, align="c",
       kind="ar"), fill=TINT, line=WINE, radius=10, anchor="m")
    B.click(g, "wipe")
    facts = tb(c, X0, 152, CW, 82, [
        Para("{k:Only before a PAST verb.} Stronger than قَد alone.", size=15,
             align="c"),
        Para("Used for news, a sudden realisation, strong feeling, or "
             "answering doubt.", size=15, align="c")], fill=TINT2, line=LINE,
        radius=10, anchor="m")
    B.click(facts)
    y = 248
    w = (CW - 24) / 3
    x = X1 - w
    for ar, tr, en in [
            ("{b:لَقَد} وَصَلْنَا أَخِيرًا", "laqad waṣalnā akhīran",
             "We've finally arrived!"),
            ("{b:لَقَد} تَغَيَّر كُلّ شَيْء", "laqad taghayyara kullu shayʾin",
             "Everything has truly changed."),
            ("{b:لَقَد} فَهِمْت الْآن", "laqad fahimtu l-āna",
             "Now I really understand.")]:
        b = exbox(c, x, y, w, 110, ar, tr, en, align="c")
        B.click(b, "wipe")
        x -= w + 12
    n = note(c, X0, 372, CW, 104, "The {t:la-} in لَقَد is the same emphatic "
             "prefix you will meet as Tool 5. Grammarians link it to an "
             "implied oath: 'I swear, it has happened.'", label="WHY IT IS STRONG",
             size=14)
    B.click(n)


@slide(2)
def p2_not_auto(s, B, ctx):
    frame(s, "Why لَقَد Is Not Automatically Right", stop=3, page=ctx.page)
    rows = [
        ("لَقَد أَذْهَب إِلَى الْعَمَل", ["أَذْهَب إِلَى الْعَمَل",
         "~adhhabu ilā l-ʿamali"], "A PAST verb is needed after لَقَد. For "
         "'may go', use قَد أَذْهَب."),
        ("لَقَد لَم أَفْهَم", ["لَم أَفْهَم", "~lam afham"],
         "No negation after لَقَد."),
        ("لَقَد سَأُسَافِر غَدًا", ["سَأُسَافِر غَدًا", "~sa-usāfiru ghadan"],
         "Future: no قَد, no لَقَد."),
    ]
    y = 66
    for wrong, right, why in rows:
        fixrow(s, B, y, wrong, right, why, h=74)
        y += 82
    n = note(s.shapes, X0, 318, CW, 160,
             "The normal report is شَرِبْت الْقَهْوَة ({t:sharibtu "
             "l-qahwata}, 'I drank the coffee'). Saying لَقَد شَرِبْتُهَا "
             "({t:laqad sharibtuhā}) sounds dramatic — natural only if someone "
             "says you didn't. Using لَقَد in every sentence weakens it.",
             label="TONE: ROUTINE INFORMATION NEEDS NO لَقَد", size=14.5)
    B.click(n)


@slide(2)
def p2_la(s, B, ctx):
    frame(s, "Tool 5 · لَـ (la-) — The Emphatic Prefix", stop=3, page=ctx.page)
    banner(s, B, "One letter with a fatha, written JOINED to the next word. "
           "It is not an independent word.")
    data = [
        ("After إِنّ — on the Predicate", "إِنّ الْحَقّ {b:لَ}وَاضِح",
         "inna l-ḥaqqa la-wāḍiḥun", "The truth is clearly evident.",
         "Name: اللَّام الْمُزَحْلَقَة ({t:al-lām al-muzaḥlaqa}, 'the moved lām')"),
        ("At the Start of a Sentence", "{b:لَ}أَنْت صَدِيق وَفِيّ",
         "la-anta ṣadīqun wafiyyun", "You truly are a loyal friend.",
         "FORMAL / literary. Name: لَام الِابْتِدَاء ({t:lām al-ibtidāʾ})"),
        ("Inside لَقَد", "{b:لَ}قَد نَجَحْنَا", "laqad najaḥnā",
         "We have truly succeeded.", "Prefix لَـ + قَد, before a past verb (Tool 4)"),
    ]
    w, gap = 272, 14
    x = X1 - w
    for t, ar, tr, en, n_ in data:
        g = card(s.shapes, x, 106, w, 250, t, [
            Para(ar, size=20, align="c", kind="ar", space_after=0),
            Para("{t:%s}" % tr, size=12, align="c", kind="en"),
            Para(en, size=12.5, align="c", kind="en", bold=True, space_after=8),
            Para(n_, size=11.5, align="c", kind="en", color=MUTED)],
            anchor="m", title_size=14)
        B.click(g, "wipe")
        x -= w + gap
    n = note(s.shapes, X0, 370, CW, 104, "Do not confuse it with لِـ ({t:li-}, "
             "'for / to') with a kasra: لِلطُّلَّاب ({t:li-ṭ-ṭullābi}, 'for "
             "the students'). Emphatic {t:la-} has a fatha and adds no "
             "'for'.", label="LOOK-ALIKE", size=14)
    B.click(n)


@slide(2)
def p2_buildup(s, B, ctx):
    frame(s, "Building Up: إِنّ … لَـ", stop=3, page=ctx.page)
    c = s.shapes
    data = [
        ("NEUTRAL", "الْخَبَر مُهِمّ", "al-khabaru muhimmun",
         "The news is important."),
        ("WITH إِنّ", "{b:إِنّ} الْخَبَر مُهِمّ", "inna l-khabara muhimmun",
         "The news is indeed important."),
        ("WITH إِنّ … لَـ", "{b:إِنّ} الْخَبَر {b:لَ}مُهِمّ",
         "inna l-khabara la-muhimmun", "The news is truly, really important."),
    ]
    w, aw = 256, 38
    x = X1 - w
    for i, (lab, ar, tr, en) in enumerate(data):
        b = exbox(c, x, 70, w, 116, ar, tr, en, align="c", label=lab,
                  fill=WHITE if i == 0 else LBLUE,
                  line=LINE if i == 0 else BLUE)
        B.click(b, "wipe")
        if i < 2:
            a = left_arrow(c, x - aw + 4, 116, aw - 8, 26)
            B.click(a)
        x -= w + aw
    rule = note(c, X0, 200, CW, 76, "The prefix لَـ goes on the predicate "
                "(or on a subject that comes late), never directly after "
                "إِنّ: {r:✗} إِنّ لَالْخَبَر مُهِمّ", label="WHERE THE PREFIX GOES",
                size=14)
    B.click(rule)
    w2 = (CW - 12) / 2
    for i, (ar, tr, en) in enumerate([
            ("{b:إِنَّك لَ}عَلَى حَقّ", "innaka la-ʿalā ḥaqqin",
             "You are absolutely right."),
            ("{b:إِنَّه لَ}شَرَف لِي", "innahu la-sharafun lī",
             "It is truly an honour for me.")]):
        b = exbox(c, X1 - w2 - i * (w2 + 12), 290, w2, 92, ar, tr, en,
                  align="c", label="WITH A PRONOUN")
        B.click(b, "wipe")
    n = note(c, X0, 396, CW, 82, "This double emphasis is typical of FORMAL "
             "speech and writing. In conversation, one tool is usually enough.",
             label="REGISTER", size=14)
    B.click(n)


@slide(2)
def p2_readcontext(s, B, ctx):
    frame(s, "Reading Context: Which Meaning of قَد?", stop=3, page=ctx.page)
    drill_header(s, y=62, labels=("Sentence", "Look at the verb", "Meaning"))
    items = [
        dict(q=ex("قَد يَنْزِل الْمَطَر الْيَوْم، فَخُذ مِظَلَّة.",
                  "qad yanzilu l-maṭaru l-yawma, fa-khudh miẓallatan.",
                  "It may rain today, so take an umbrella."),
             cue=["يَنْزِل", "present verb"], a="{g:✓ Possibility}\n'may'"),
        dict(q=ex("قَد انْتَهَى الِاجْتِمَاع، تَفَضَّل.",
                  "qadi ntahā l-ijtimāʿu, tafaḍḍal.",
                  "The meeting has ended — please come in."),
             cue=["انْتَهَى", "past verb"], a="{g:✓ Completion}\n'has ended'"),
        dict(q=ex("قَد يَكُون الْمُدِير فِي اجْتِمَاع.",
                  "qad yakūnu l-mudīru fī jtimāʿin.",
                  "The manager may be in a meeting."),
             cue=["يَكُون", "present verb"], a="{g:✓ Possibility}\n'may be'"),
        dict(q=ex("وَقَد أَعْلَنَت الْوِزَارَة النَّتَائِج.",
                  "wa-qad aʿlanati l-wizāratu n-natāʾija.",
                  "The ministry has announced the results."),
             cue=["أَعْلَنَت", "past verb"],
             a="{g:✓ Completion}\nFORMAL news style"),
    ]
    drill(s, B, items, y0=86, rowh=92, gap=8)


@slide(2)
def p2_convnews(s, B, ctx):
    frame(s, "In Conversation and in the News", stop=3, page=ctx.page)
    tb(s.shapes, X0, 62, CW, 22, Para("EVERYDAY · Two colleagues", size=12.5,
       color=MUTED, align="r"), margins=(0, 0, 0, 0))
    lines = [
        ("r", "مُنَى · Muna", "هَل وَصَلَت رِسَالَتِي؟", "hal waṣalat risālatī?",
         "Did my message arrive?"),
        ("l", "خَالِد · Khālid", "نَعَم، {b:لَقَد} وَصَلَت وَقَرَأْتُهَا.",
         "naʿam, laqad waṣalat wa-qaraʾtuhā.",
         "Yes, it did arrive, and I've read it."),
        ("r", "مُنَى · Muna", "وَمَتَى سَتَرُدّ؟", "wa-matā sa-taruddu?",
         "And when will you reply?"),
        ("l", "خَالِد · Khālid", "{b:قَد} أَرُدّ اللَّيْلَة.",
         "qad aruddu l-laylata.", "I may reply tonight."),
    ]
    y = 86
    for side, name, ar, tr, en in lines:
        bubble(s, B, y, side, name, ar, tr, en, h=72, w=560)
        y += 78
    nb = exbox(s.shapes, X0, 404, CW, 76,
               "وَ{b:قَد} أَكَّدَت الشَّرِكَة أَنّ الْخِدْمَة سَتَعُود قَرِيبًا.",
               "wa-qad akkadati sh-sharikatu anna l-khidmata sa-taʿūdu qarīban.",
               "The company has confirmed that the service will return soon.",
               label="FORMAL · NEWS REPORT", fill=TINT, line=WINE)
    B.click(nb, "wipe")


@slide(2)
def p2_table(s, B, ctx):
    frame(s, "Verb-Form Comparison Table", stop=3, page=ctx.page)
    rows = [
        ("ذَهَب", "dhahaba", "He went.", "neutral past"),
        ("قَد ذَهَب", "qad dhahaba", "He has (already) gone.", "completion"),
        ("لَقَد ذَهَب", "laqad dhahaba", "He has indeed gone.",
         "completion, stressed"),
        ("يَذْهَب", "yadhhabu", "He goes / is going.", "neutral present"),
        ("قَد يَذْهَب", "qad yadhhabu", "He may go.", "possibility"),
        ("قَد لَا يَذْهَب", "qad lā yadhhabu", "He may not go.",
         "negative possibility"),
        ("سَيَذْهَب", "sa-yadhhabu", "He will go.", "future — no قَد"),
    ]
    rows = [(a, "{t:%s}" % b, c_, d) for a, b, c_, d in rows]
    table(s, B, X0, 62, [190, 190, 250, 214], ["Arabic", "Sound",
          "Translation", "What It Shows"], rows, rowh=42, ar_size=18,
          en_size=12)
    w = 416
    a = note(s.shapes, X1 - w, 392, w, 98, "These combinations are not "
             "used: {r:✗} لَقَد يَذْهَب · {r:✗} قَد سَيَذْهَب", label="NOT USED",
             size=14, color=RED, fill=LRED)
    b = note(s.shapes, X0, 392, w, 98, "Further ahead: لَأُسَاعِدَنَّك "
             "({t:la-usāʿidannaka}, 'I will certainly help you') uses an "
             "emphatic {t:-nna} ending — not one of this chapter's tools.",
             label="YOU MAY ALSO MEET", size=12.5)
    B.click(a)
    B.click(b)


@slide(2)
def p2_practice(s, B, ctx):
    frame(s, "Practice: قَد, لَقَد, لَـ — or Nothing?", stop=3, page=ctx.page)
    drill_header(s, y=62, labels=("Sentence", "Cue", "Answer"))
    B_ = "ــــــ"
    items = [
        dict(q=[f"{B_} يَتَأَخَّر الْبَاص بِسَبَب الزَّحْمَة",
                "The bus may be late because of traffic."],
             cue="present verb + 'may'",
             a=["{g:قَد}", "~qad yataʾakhkharu l-bāṣu"]),
        dict(q=[f"{B_} نَجَحْت فِي الِامْتِحَان!", "I've actually passed the exam!"],
             cue="past verb + strong feeling",
             a=["{g:لَقَد}", "~laqad najaḥtu fī l-imtiḥāni"]),
        dict(q=[f"{B_} سَأُسَافِر غَدًا", "I will travel tomorrow."],
             cue="future verb",
             a=["{g:nothing}", "~sa-usāfiru ghadan"]),
        dict(q=[f"{B_} أَنْهَيْنَا الْمَشْرُوع", "We have completed the project."],
             cue="past verb, completion",
             a=["{g:قَد}", Para("or لَقَد for more force", size=11.5,
                                 kind="en"), "~qad anhaynā l-mashrūʿa"]),
        dict(q=[f"إِنّ النَّتِيجَة {B_}رَائِعَة", "The result is truly wonderful."],
             cue="predicate after إِنّ",
             a=["{g:لَـ} ← لَرَائِعَة", "~inna n-natījata la-rāʾiʿatun"]),
        dict(q=[f"{B_} أَعْمَل فِي الْمَكْتَب كُلّ يَوْم",
                "I work in the office every day."],
             cue="routine fact",
             a=["{g:nothing}", "~aʿmalu fī l-maktabi kulla yawmin"]),
    ]
    drill(s, B, items, y0=86, rowh=64, gap=5)


# ===================================================================== PART 3
@slide(3)
def p3_open(s, B, ctx):
    opener(s, B, ctx, 4, "Emphasis, Restriction and Completeness",
           "Limiting a statement · stressing identity · stressing totality",
           "الْقَصْر وَالتَّوْكِيد الْمَعْنَوِيّ", "al-qaṣr wa-t-tawkīd al-maʿnawī",
           "Use إِنَّمَا to limit a statement, and the nouns نَفْس، عَيْن، "
           "كُلّ، جَمِيع with the right pronoun suffix.",
           [("only, merely", "إِنَّمَا", "innamā"),
            ("self, same", "نَفْس", "nafs"),
            ("the very same", "عَيْن", "ʿayn"),
            ("all, every", "كُلّ", "kull"),
            ("all, altogether", "جَمِيع", "jamīʿ")],
           ["Use إِنَّمَا to restrict a statement, and know what it does NOT mean.",
            "Make نَفْس and عَيْن agree with the noun they follow.",
            "Choose correctly between كُلّ and جَمِيع."])


@slide(3)
def p3_innama(s, B, ctx):
    frame(s, "Tool 6 · إِنَّمَا (innamā) — Only, Merely", stop=4, page=ctx.page)
    c = s.shapes
    g = group(c)
    tb(g.shapes, 700, 66, 120, 64, Para("إِنّ", size=24, align="c", kind="ar"),
       fill=WHITE, line=WINE, radius=10, anchor="m")
    tb(g.shapes, 660, 80, 40, 36, Para("+", size=22, align="c"), anchor="m")
    tb(g.shapes, 540, 66, 120, 64, [Para("مَا", size=24, align="c", kind="ar")],
       fill=WHITE, line=WINE, radius=10, anchor="m")
    tb(g.shapes, 500, 80, 40, 36, Para("=", size=22, align="c"), anchor="m")
    tb(g.shapes, 340, 66, 160, 64, Para("{b:إِنَّمَا}", size=26, align="c",
       kind="ar"), fill=LBLUE, line=BLUE, radius=10, anchor="m")
    B.click(g, "wipe")
    facts = tb(c, X0, 142, CW, 100, [
        Para("{k:Meaning:} restriction ({t:qaṣr}) — 'only, merely, nothing "
             "more than'. It limits; it does not just strengthen.", size=14),
        Para("{k:Grammar:} the 'restraining {t:mā}' cancels إِنّ's effect, "
             "so endings stay normal. It can come before a noun or a verb.",
             size=14),
        Para("{k:Register:} FORMAL and written; everyday speech prefers "
             "فَقَط ({t:faqaṭ}, 'only').", size=14)], fill=TINT2, line=LINE,
        radius=10, anchor="m", margins=(12, 4, 12, 4))
    B.click(facts)
    w = (CW - 12) / 2
    a = exbox(c, X1 - w, 256, w, 110,
              "{b:إِنَّمَا} الْأَعْمَال بِالنِّيَّات",
              "innamā l-aʿmālu bi-n-niyyāti",
              "Actions are judged only by intentions.", align="c",
              label="FAMOUS HADITH · BEFORE A NOUN")
    b = exbox(c, X0, 256, w, 110, "{b:إِنَّمَا} أَرَدْت الْمُسَاعَدَة",
              "innamā aradtu l-musāʿadata", "I only wanted to help.",
              align="c", label="BEFORE A VERB")
    B.click(a, "wipe")
    B.click(b, "wipe")
    n = note(c, X0, 380, CW, 96, "After إِنَّمَا the subject keeps {t:-u}: "
             "{t:innamā l-aʿmāl}{k:u} — compare {t:inna l-aʿmāl}{k:a}.",
             label="NO CASE CHANGE", size=15)
    B.click(n)


@slide(3)
def p3_restrict(s, B, ctx):
    frame(s, "How Restriction Works with إِنَّمَا", stop=4, page=ctx.page)
    banner(s, B, "The limited idea usually comes LAST. Everything else is "
           "excluded.")
    rows = [
        ("{b:إِنَّمَا} أَنَا مُوَظَّف جَدِيد\n~innamā ana muwaẓẓafun jadīdun",
         "I'm only a new employee.", "what I am: nothing more"),
        ("{b:إِنَّمَا} الْعِلْم بِالتَّعَلُّم\n~innamā l-ʿilmu bi-t-taʿallumi",
         "Knowledge comes only through learning.", "the way to knowledge: one way only"),
        ("لَم أَقْصِد الْإِسَاءَة، {b:وَإِنَّمَا} أَرَدْت الْمُسَاعَدَة\n~lam aqṣidi l-isāʾata, wa-innamā aradtu l-musāʿadata",
         "I didn't mean to offend; I only wanted to help.",
         "after a negative, وَإِنَّمَا = 'but rather'"),
    ]
    table(s, B, X0, 104, [400, 250, 194], ["Example", "Translation",
          "What Is Limited"], rows, rowh=78, ar_size=18, en_size=12)
    w = (CW - 12) / 2
    a = exbox(s.shapes, X1 - w, 380, w, 104,
              "مَا أَنَا إِلَّا مُوَظَّف جَدِيد",
              "mā ana illā muwaẓẓafun jadīdun", "I'm nothing but a new employee.",
              align="c", label="ANOTHER RESTRICTION PATTERN (often stronger)")
    b = exbox(s.shapes, X0, 380, w, 104, "أَنَا مُوَظَّف جَدِيد فَقَط",
              "ana muwaẓẓafun jadīdun faqaṭ", "I'm only a new employee.",
              align="c", label="EVERYDAY: فَقَط 'only'")
    B.click(a, "wipe")
    B.click(b, "wipe")


@slide(3)
def p3_nafs(s, B, ctx):
    frame(s, "Tool 7 · نَفْس (nafs) — Self, Same", stop=4, page=ctx.page)
    banner(s, B, "The word نَفْس is a NOUN (literally 'self, soul'), not a "
           "particle. Its position changes the meaning.")
    data = [
        ("AFTER a Noun + Suffix", "'-self' (emphasis)",
         "حَضَر الْمُدِير {b:نَفْسُه}", "ḥaḍara l-mudīru nafsuhu",
         "The manager himself attended.",
         "Follows the noun's case. Can also mean 'the same' in formal style."),
        ("BEFORE a Noun", "'the same'",
         "نَسْكُن فِي {b:نَفْس} الْحَيّ", "naskunu fī nafsi l-ḥayyi",
         "We live in the same neighbourhood.",
         "A possessive structure, very common in modern Arabic."),
        ("With the Prefix بِـ", "'personally, by myself'",
         "سَأَتَكَلَّم مَعَه {b:بِنَفْسِي}", "sa-atakallamu maʿahu bi-nafsī",
         "I'll speak to him myself.", "بِنَفْسِي · بِنَفْسِك · بِنَفْسِه …"),
    ]
    w, gap = 272, 14
    x = X1 - w
    for t, m, ar, tr, en, n_ in data:
        g = card(s.shapes, x, 106, w, 262, t, [
            Para(m, size=13, color=BLUE, bold=True, align="c", kind="en",
                 space_after=6),
            Para(ar, size=19, align="c", kind="ar", space_after=0),
            Para("{t:%s}" % tr, size=12, align="c", kind="en"),
            Para(en, size=12.5, align="c", kind="en", bold=True, space_after=8),
            Para(n_, size=11.5, align="c", kind="en", color=MUTED)],
            anchor="m", title_size=14.5)
        B.click(g, "wipe")
        x -= w + gap
    n = note(s.shapes, X0, 382, CW, 94, "Only the first use is grammatical "
             "emphasis ({t:tawkīd}). نَفْس الشَّخْص ({t:nafsu sh-shakhṣi}) = "
             "'the same person'; الشَّخْص نَفْسُه ({t:ash-shakhṣu nafsuhu}) = "
             "'the person himself' (or, formally, 'the same person').",
             label="KEEP THEM APART", size=13.5)
    B.click(n)


@slide(3)
def p3_ayn(s, B, ctx):
    frame(s, "Tool 8 · عَيْن (ʿayn) — The Very Same", stop=4, page=ctx.page)
    banner(s, B, "The word عَيْن is a NOUN (literally 'eye'). After a noun "
           "with a suffix it means 'the very same'. FORMAL.")
    rows = [
        ("هَذَا هُوَ الْكِتَاب {b:عَيْنُه}\n~hādhā huwa l-kitābu ʿaynuhu",
         "This is the very same book.", "subject case: {t:ʿaynu-hu}"),
        ("رَأَيْت الرَّجُل {b:عَيْنَه}\n~raʾaytu r-rajula ʿaynahu",
         "I saw the very same man.", "object case: {t:ʿayna-hu}"),
        ("الْفِكْرَة {b:عَيْنُهَا}\n~al-fikratu ʿaynuhā",
         "the very same idea", "feminine noun: {t:-hā}"),
        ("هَذَا {b:عَيْن} الصَّوَاب\n~hādhā ʿaynu ṣ-ṣawābi",
         "This is exactly right.", "fixed expression, before a noun"),
    ]
    table(s, B, X0, 104, [360, 260, 224], ["Example", "Translation",
          "Agreement"], rows, rowh=68, ar_size=18, en_size=12)
    n = note(s.shapes, X0, 410, CW, 70, "Like نَفْس, the word عَيْن follows "
             "the case of its noun and takes a matching pronoun suffix. In "
             "everyday speech, نَفْس is far more common.", label="AGREEMENT",
             size=14)
    B.click(n)


@slide(3)
def p3_agree(s, B, ctx):
    frame(s, "Agreement: The Pronoun Suffix Must Match", stop=4, page=ctx.page)
    rows = [
        ("masculine singular\n(الْمُدِير)", "الْمُدِير نَفْسُه\n~al-mudīru nafsuhu",
         "الْكِتَاب كُلُّه\n~al-kitābu kulluhu", "{t:-hu}"),
        ("feminine singular\n(الْمُدِيرَة)", "الْمُدِيرَة نَفْسُهَا\n~al-mudīratu nafsuhā",
         "الْمَدِينَة كُلُّهَا\n~al-madīnatu kulluhā", "{t:-hā}"),
        ("masculine plural\n(people)", "الطُّلَّاب أَنْفُسُهُم\n~aṭ-ṭullābu anfusuhum",
         "الطُّلَّاب كُلُّهُم\n~aṭ-ṭullābu kulluhum", "{t:-hum}"),
        ("feminine plural\n(people)", "الطَّالِبَات أَنْفُسُهُنّ\n~aṭ-ṭālibātu anfusuhunna",
         "الطَّالِبَات كُلُّهُنّ\n~aṭ-ṭālibātu kulluhunna", "{t:-hunna}"),
        ("plural of things", "—", "الْكُتُب كُلُّهَا\n~al-kutubu kulluhā",
         "{t:-hā} (things)"),
    ]
    table(s, B, X0, 62, [200, 250, 250, 144], ["The Noun", "with نَفْس",
          "with كُلّ", "Suffix"], rows, rowh=62, ar_size=18, en_size=12)
    n = note(s.shapes, X0, 406, CW, 84, "The word also copies the case: "
             "جَاء الْمُدِير نَفْسُه ({t:nafsu-hu}) · رَأَيْت الْمُدِير نَفْسَه "
             "({t:nafsa-hu}) · سَلَّمْت عَلَى الْمُدِير نَفْسِه ({t:nafsi-hu}). "
             "Plural of نَفْس: أَنْفُس ({t:anfus}).", label="CASE TOO", size=13)
    B.click(n)


@slide(3)
def p3_kull(s, B, ctx):
    frame(s, "Tool 9 · كُلّ (kull) — All, Every, Whole", stop=4, page=ctx.page)
    rows = [
        ("كُلّ + indefinite singular", "كُلّ يَوْم · كُلّ طَالِب\n~kulla yawmin · kullu ṭālibin",
         "every day · every student", "quantifier"),
        ("كُلّ + definite plural", "كُلّ الطُّلَّاب\n~kullu ṭ-ṭullābi",
         "all the students", "quantifier"),
        ("كُلّ + definite singular", "كُلّ الْيَوْم\n~kulla l-yawmi",
         "the whole day", "quantifier"),
        ("noun + كُلّ + suffix", "نَجَح الطُّلَّاب {b:كُلُّهُم}\n~najaḥa ṭ-ṭullābu kulluhum",
         "The students all passed.", "{b:emphasis (tawkīd)}"),
        ("noun + كُلّ + suffix", "قَرَأْت الْكِتَاب {b:كُلَّه}\n~qaraʾtu l-kitāba kullahu",
         "I read the entire book.", "{b:emphasis (tawkīd)}"),
    ]
    table(s, B, X0, 62, [210, 280, 200, 154], ["Pattern", "Example",
          "Meaning", "Job"], rows, rowh=62, ar_size=18, en_size=12,
          fills={(3, 3): LBLUE, (4, 3): LBLUE})
    n = note(s.shapes, X0, 406, CW, 84, "Only after the noun is كُلّ "
             "grammatical emphasis. Before a noun it is an ordinary "
             "quantifier ('every', 'all', 'whole').", label="GRAMMAR NOTE",
             size=14)
    B.click(n)


@slide(3)
def p3_jami(s, B, ctx):
    frame(s, "Tool 10 · جَمِيع (jamīʿ) — All, Altogether", stop=4, page=ctx.page)
    data = [
        ("Before a Definite Plural", "{b:جَمِيع} الْمُوَظَّفِين",
         "jamīʿu l-muwaẓẓafīna", "all the employees", "Common in FORMAL notices."),
        ("After a Noun + Suffix", "جَاء الضُّيُوف {b:جَمِيعُهُم}",
         "jāʾa ḍ-ḍuyūfu jamīʿuhum", "The guests came, all of them.",
         "Grammatical emphasis, like كُلُّهُم."),
        ("'All Together': جَمِيعًا", "شُكْرًا لَكُم {b:جَمِيعًا}",
         "shukran lakum jamīʿan", "Thank you all.",
         "Very common EVERYDAY phrase."),
    ]
    w, gap = 272, 14
    x = X1 - w
    for t, ar, tr, en, n_ in data:
        g = card(s.shapes, x, 66, w, 230, t, [
            Para(ar, size=20, align="c", kind="ar", space_after=0),
            Para("{t:%s}" % tr, size=12, align="c", kind="en"),
            Para(en, size=12.5, align="c", kind="en", bold=True, space_after=8),
            Para(n_, size=12, align="c", kind="en", color=MUTED)],
            anchor="m", title_size=14.5)
        B.click(g, "wipe")
        x -= w + gap
    fixrow(s, B, 316, "جَمِيع يَوْم", ["كُلّ يَوْم", "~kulla yawmin — every day"],
           "Never use جَمِيع for 'every' with a single indefinite noun.", h=76)
    n = note(s.shapes, X0, 404, CW, 76, "Also everyday: ذَهَبْنَا جَمِيعًا "
             "({t:dhahabnā jamīʿan}) = 'we all went (together)'.",
             label="ALSO", size=14)
    B.click(n)


@slide(3)
def p3_kull_vs_jami(s, B, ctx):
    frame(s, "Choosing Between كُلّ and جَمِيع", stop=4, page=ctx.page)
    G = lambda t: ("{g:✓} " + t, LGREEN)
    R = lambda t: ("{r:✗} " + t, LRED)
    rows = [
        ("every (+ one indefinite noun)", G("كُلّ طَالِب"), R("not used")),
        ("all (+ definite plural)", G("كُلّ الطُّلَّاب"), G("جَمِيع الطُّلَّاب")),
        ("the whole (+ definite singular)", G("كُلّ الْيَوْم"),
         R("not usual")),
        ("emphasis after a noun", G("الطُّلَّاب كُلُّهُم"),
         G("الطُّلَّاب جَمِيعُهُم")),
        ("'all together' (adverb)", R("—"), G("جَمِيعًا")),
    ]
    table(s, B, X0, 62, [284, 280, 280], ["Meaning", "كُلّ (kull)",
          "جَمِيع (jamīʿ)"], rows, rowh=58, ar_size=18, en_size=12.5)
    n = note(s.shapes, X0, 386, CW, 96, "The all-rounder is كُلّ. The word "
             "جَمِيع means 'all' of a group and sounds more FORMAL — except "
             "the everyday جَمِيعًا. Avoid جَمِيع for 'every' or 'the whole'.",
             label="RULE OF THUMB", size=14)
    B.click(n)
    legend(s, [(GREEN, LGREEN, "✓ green = used"), (RED, LRED, "✗ red = not used")],
           y=484)


@slide(3)
def p3_chart(s, B, ctx):
    frame(s, "What Each Construction Emphasizes — and What It Does Not Mean",
          stop=4, page=ctx.page)
    rows = [
        ("إِنّ", "the truth of the whole statement", "not 'if' (إِنْ); not 'only'"),
        ("أَنّ", "a fact reported inside a sentence ('that')", "not a sentence opener; not أَنْ + verb"),
        ("قَد + past", "completion, certainty", "not possibility"),
        ("قَد + present", "possibility, 'may'", "not certainty (modern usage)"),
        ("لَقَد", "a completed event, stated strongly", "never with present, future or negation"),
        ("لَـ", "reinforcement on one word", "not a separate word; not لِـ 'for'"),
        ("إِنَّمَا", "restriction: 'only'", "not simple 'indeed'"),
        ("نَفْس / عَيْن", "identity: '-self', 'the very same'", "not 'all'"),
        ("كُلّ / جَمِيع", "totality: 'all of it / them'", "not 'the same'"),
    ]
    table(s, B, X0, 62, [150, 360, 334], ["Tool", "Emphasizes",
          "Does NOT Mean"], rows, rowh=42, ar_size=18, en_size=12.5,
          fills={(i, 2): LRED for i in range(9)})
    legend(s, [(RED, LRED, "red column = what the tool does NOT mean")], y=482)


@slide(3)
def p3_fourwords(s, B, ctx):
    frame(s, "Four English Words, Four Different Arabic Jobs", stop=4,
          page=ctx.page)
    banner(s, B, "English words do not map one-to-one. Choose by the JOB, "
           "not by the English word.")
    data = [
        ("'indeed'", "strengthen the truth", "إِنّ · لَقَد",
         "{b:لَقَد} فَازَ فَرِيقُنَا", "laqad fāza farīqunā",
         "Our team has indeed won."),
        ("'only'", "limit the statement", "إِنَّمَا · فَقَط",
         "{b:إِنَّمَا} أَسْأَل", "innamā asʾalu", "I'm only asking."),
        ("'himself / the same'", "stress identity", "نَفْس · عَيْن",
         "قَابَلْت الْكَاتِب {b:نَفْسَه}", "qābaltu l-kātiba nafsahu",
         "I met the writer himself."),
        ("'all'", "stress totality", "كُلّ · جَمِيع",
         "وَصَل الْفَرِيق {b:كُلُّه}", "waṣala l-farīqu kulluhu",
         "The whole team arrived."),
    ]
    w, gap = 202, 12
    x = X1 - w
    for t, job, tools, ar, tr, en in data:
        g = card(s.shapes, x, 106, w, 300, t, [
            Para(job.upper(), size=10.5, color=MUTED, bold=True, align="c",
                 kind="en"),
            Para(tools, size=20, align="c", kind="ar", space_after=12),
            Para(ar, size=18, align="c", kind="ar", space_after=0),
            Para("{t:%s}" % tr, size=11.5, align="c", kind="en"),
            Para(en, size=12, align="c", kind="en", bold=True)],
            anchor="m", title_size=15)
        B.click(g, "wipe")
        x -= w + gap
    n = note(s.shapes, X0, 418, CW, 60, "English 'indeed' is often best left "
             "out of a translation of إِنّ — use stress or 'really' instead.",
             label="TRANSLATION TIP", size=14)
    B.click(n)


@slide(3)
def p3_mistakes(s, B, ctx):
    frame(s, "Common Mistakes: Restriction, Identity and Totality", stop=4,
          page=ctx.page)
    rows = [
        ("حَضَر الْمُدِير نَفْس", ["حَضَر الْمُدِير نَفْسُه",
         "~ḥaḍara l-mudīru nafsuhu"], "'-self' needs a pronoun suffix."),
        ("رَأَيْت الطَّالِبَة نَفْسَه", ["رَأَيْت الطَّالِبَة نَفْسَهَا",
         "~raʾaytu ṭ-ṭālibata nafsahā"], "The suffix agrees: feminine takes {t:-hā}."),
        ("نَسْكُن فِي الْحَيّ نَفْس", ["نَسْكُن فِي نَفْس الْحَيّ",
         "~naskunu fī nafsi l-ḥayyi"], "'the same' + noun: نَفْس comes first."),
        ("قَرَأْت الْكُتُب كُلَّهُم", ["قَرَأْت الْكُتُب كُلَّهَا",
         "~qaraʾtu l-kutuba kullahā"], "Plural of things takes {t:-hā}."),
        ("إِنَّمَا الْمُشَارِكِين طُلَّاب", ["إِنَّمَا الْمُشَارِكُون طُلَّاب",
         "~innamā l-mushārikūna ṭullābun"], "No case change after إِنَّمَا."),
    ]
    y = 64
    for wrong, right, why in rows:
        fixrow(s, B, y, wrong, right, why, h=78)
        y += 84


@slide(3)
def p3_practice(s, B, ctx):
    frame(s, "Practice: Complete with the Right Word", stop=4, page=ctx.page)
    drill_header(s, y=62, labels=("Sentence", "Meaning needed", "Answer"))
    B_ = "ــــــ"
    items = [
        dict(q=[f"سَأُكَلِّم الْمُدِير {B_}", "I'll speak to the manager himself."],
             cue="'himself' (object)",
             a=["{g:نَفْسَه}", "~sa-ukallimu l-mudīra nafsahu"]),
        dict(q=[f"{B_} أَنَا زَائِر، لَسْت مِن هُنَا", "I'm only a visitor; I'm not from here."],
             cue="'only'",
             a=["{g:إِنَّمَا}", "~innamā ana zāʾirun"]),
        dict(q=[f"حَضَر الطُّلَّاب {B_}", "The students all attended."],
             cue="'all of them'",
             a=["{g:كُلُّهُم} / {g:جَمِيعُهُم}", "~kulluhum / jamīʿuhum"]),
        dict(q=[f"نَعْمَل فِي {B_} الشَّرِكَة", "We work in the same company."],
             cue="'the same' + noun",
             a=["{g:نَفْس}", "~naʿmalu fī nafsi sh-sharikati"]),
        dict(q=[f"شُكْرًا لَكُم {B_}", "Thank you all."],
             cue="'all (together)'",
             a=["{g:جَمِيعًا}", "~shukran lakum jamīʿan"]),
        dict(q=[f"هَذَا هُوَ الْعُنْوَان {B_}", "This is the very same address."],
             cue="'the very same' (FORMAL)",
             a=["{g:عَيْنُه}", "~hādhā huwa l-ʿunwānu ʿaynuhu"]),
    ]
    drill(s, B, items, y0=86, rowh=64, gap=5)


# ===================================================================== PART 4
@slide(4)
def p4_open(s, B, ctx):
    opener(s, B, ctx, 5, "Putting the Tools Together",
           "Choosing the right tool for the meaning you intend",
           "اخْتِيَار الْأَدَاة الْمُنَاسِبَة", "ikhtiyār al-adāti l-munāsiba",
           "Work out what you want to stress, check the structure, then "
           "pick the tool — in conversation, at work and in formal Arabic.",
           [("strengthen a statement", "إِنّ · أَنّ", "inna · anna"),
            ("complete / possible", "قَد · لَقَد", "qad · laqad"),
            ("reinforce a word", "لَـ", "la-"),
            ("limit", "إِنَّمَا", "innamā"),
            ("identity & totality", "نَفْس · كُلّ", "nafs · kull …")],
           ["Identify the PURPOSE of emphasis in any sentence.",
            "Apply the three-step method: meaning, then structure, then tool.",
            "Produce natural emphatic sentences for real situations."])


@slide(4)
def p4_master(s, B, ctx):
    frame(s, "Master Map of the Chapter", stop=5, page=ctx.page)
    tools = [
        ("1", "إِنّ", "inna", "strengthens a whole nominal sentence"),
        ("2", "أَنّ", "anna", "'that' — a fact inside a sentence"),
        ("3", "قَد", "qad", "past: already · present: may"),
        ("4", "لَقَد", "laqad", "strong completion (past only)"),
        ("5", "لَـ", "la-", "prefix: reinforces one word"),
        ("6", "إِنَّمَا", "innamā", "restriction: only"),
        ("7", "نَفْس", "nafs", "-self; 'the same' before a noun"),
        ("8", "عَيْن", "ʿayn", "the very same (FORMAL)"),
        ("9", "كُلّ", "kull", "all, every, whole"),
        ("10", "جَمِيع", "jamīʿ", "all; جَمِيعًا all together"),
    ]
    zone = {"1": BLUE, "2": BLUE, "6": BLUE, "3": WINE, "4": WINE, "5": GREEN,
            "7": "7A5C00", "8": "7A5C00", "9": "7A5C00", "10": "7A5C00"}
    w, h, gx, gy = 160, 150, 11, 12
    tiles = []
    for i, (n, ar, tr, en) in enumerate(tools):
        r_, k = divmod(i, 5)
        x = X1 - w - k * (w + gx)
        y = 66 + r_ * (h + gy)
        col = zone[n]
        tiles.append(tb(s.shapes, x, y, w, h, [
            Para(f"TOOL {n}", size=10.5, color=col, bold=True, align="c",
                 kind="en"),
            Para(ar, size=24, align="c", kind="ar", space_after=0),
            Para("{t:%s}" % tr, size=12, align="c", kind="en"),
            Para(en, size=11.5, align="c", kind="en")], fill=WHITE, line=col,
            lw=1.5, radius=10, anchor="m", margins=(6, 4, 6, 4)))
    B.chain(tiles[:5])
    B.chain(tiles[5:])
    legend(s, [(BLUE, WHITE, "sentence particles"),
               (WINE, WHITE, "verb particles"),
               (GREEN, WHITE, "attached prefix"),
               ("7A5C00", WHITE, "nouns used for emphasis")], y=402)
    n = note(s.shapes, X0, 430, CW, 56, "Numbers follow the teaching order of "
             "this chapter.", label="NOTE", size=13)
    B.click(n)


@slide(4)
def p4_purpose(s, B, ctx):
    frame(s, "What Is the Speaker Trying to Do?", stop=5, page=ctx.page)
    rows = [
        ("Insist that a statement is true?", "إِنّ · إِنّ … لَـ",
         "{b:إِنّ} الْخُطَّة نَاجِحَة\n~inna l-khuṭṭata nājiḥatun"),
        ("Report a fact inside a sentence?", "أَنّ",
         "أَعْرِف {b:أَنّ} الْخُطَّة نَاجِحَة\n~aʿrifu anna l-khuṭṭata nājiḥatun"),
        ("Present an action as really done?", "قَد / لَقَد + past",
         "{b:لَقَد} نَجَحَت الْخُطَّة\n~laqad najaḥati l-khuṭṭatu"),
        ("Say something may happen?", "قَد + present",
         "{b:قَد} تَنْجَح الْخُطَّة\n~qad tanjaḥu l-khuṭṭatu"),
        ("Limit the statement?", "إِنَّمَا",
         "{b:إِنَّمَا} هِيَ فِكْرَة\n~innamā hiya fikratun"),
        ("Stress who or how much?", "نَفْس · عَيْن · كُلّ · جَمِيع",
         "الْفَرِيق {b:كُلُّه} مُوَافِق\n~al-farīqu kulluhu muwāfiqun"),
    ]
    table(s, B, X0, 62, [290, 220, 334], ["The Speaker Wants To…", "Tool",
          "Example"], rows, rowh=64, ar_size=17.5, en_size=12.5)
    tb(s.shapes, X0, 476, CW, 24, Para("Translations: The plan is working. · I "
       "know the plan is working. · The plan has succeeded. · The plan may "
       "succeed. · It's only an idea. · The whole team agrees.", size=10.5,
       color=MUTED, align="c"), margins=(0, 0, 0, 0))


@slide(4)
def p4_method(s, B, ctx):
    frame(s, "A Three-Step Selection Method", stop=5, page=ctx.page)
    steps = [
        ("STEP 1 · Meaning", "What do I want to stress?",
         "identity: 'the manager HIMSELF'"),
        ("STEP 2 · Structure", "What is the word or sentence type?",
         "a noun: الْمُدِير, masculine singular, subject"),
        ("STEP 3 · Tool", "Which construction fits?",
         "use نَفْس + {t:-hu} (nominative): نَفْسُه"),
    ]
    w, aw = 256, 38
    x = X1 - w
    for i, (t, q, a) in enumerate(steps):
        g = card(s.shapes, x, 70, w, 210, t, [
            Para(q, size=15, align="c", kind="en", bold=True, space_after=10),
            Para("Example:", size=11, color=MUTED, align="c", kind="en"),
            Para(a, size=13, align="c", color=BLUE)], anchor="m", title_size=15)
        B.click(g, "wipe")
        if i < 2:
            ar_ = left_arrow(s.shapes, x - aw + 4, 160, aw - 8, 28)
            B.click(ar_)
        x -= w + aw
    res = exbox(s.shapes, X0 + 140, 300, CW - 280, 104,
                "رَدّ الْمُدِير {b:نَفْسُه} عَلَى رِسَالَتِي",
                "radda l-mudīru nafsuhu ʿalā risālatī",
                "The manager himself replied to my message.", align="c",
                label="RESULT", fill=LGREEN, line=GREEN, label_color=GREEN)
    B.click(res, "wipe")
    n = note(s.shapes, X0, 418, CW, 64, "If no emphasis is needed, the best "
             "choice is often NO tool.", label="STEP 0", size=14)
    B.click(n)


@slide(4)
def p4_worked(s, B, ctx):
    frame(s, "Worked Examples with the Method", stop=5, page=ctx.page)
    rows = [
        ("We've finally finished!", "completed event + feeling",
         "past verb", "{b:لَقَد} انْتَهَيْنَا أَخِيرًا\n~laqadi ntahaynā akhīran"),
        ("It may rain this evening.", "possibility", "present verb",
         "{b:قَد} تُمْطِر مَسَاءً\n~qad tumṭiru masāʾan"),
        ("I'm only asking.", "restriction", "verb", "{b:إِنَّمَا} أَسْأَل\n~innamā asʾalu"),
        ("Arabic is truly beautiful.", "strengthen a statement",
         "nominal sentence", "{b:إِنّ} الْعَرَبِيَّة جَمِيلَة\n~inna l-ʿarabiyyata jamīlatun"),
        ("I read all the messages.", "totality", "plural of things",
         "قَرَأْت الرَّسَائِل {b:كُلَّهَا}\n~qaraʾtu r-rasāʾila kullahā"),
    ]
    table(s, B, X0, 62, [220, 190, 150, 284], ["I Want to Say", "1 · Meaning",
          "2 · Structure", "3 · Arabic"], rows, rowh=72, ar_size=18,
          en_size=12.5)
    legend(s, [(BLUE, LBLUE, "blue = the selected tool")], y=482)


@slide(4)
def p4_conv(s, B, ctx):
    frame(s, "Everyday Conversation: Planning a Trip", stop=5, page=ctx.page)
    lines = [
        ("r", "لَيْلَى · Layla", "هَل حَجَزْت الْفُنْدُق؟",
         "hal ḥajazta l-funduqa?", "Have you booked the hotel?"),
        ("l", "سَامِي · Sāmī", "نَعَم، {b:لَقَد} حَجَزْتُه أَمْس.",
         "naʿam, laqad ḥajaztuhu amsi.", "Yes — I booked it yesterday. It's done."),
        ("r", "لَيْلَى · Layla", "وَالتَّذَاكِر؟", "wa-t-tadhākiru?",
         "And the tickets?"),
        ("l", "سَامِي · Sāmī", "{b:قَد} أَشْتَرِيهَا اللَّيْلَة؛ {b:إِنَّمَا} أَنْتَظِر سِعْرًا أَفْضَل.",
         "qad ashtarīhā l-laylata; innamā antaẓiru siʿran afḍala.",
         "I may buy them tonight — I'm only waiting for a better price."),
        ("r", "لَيْلَى · Layla", "مُمْتَاز! سَنُسَافِر {b:كُلُّنَا} مَعًا.",
         "mumtāzun! sa-nusāfiru kullunā maʿan.",
         "Excellent! We'll all travel together."),
    ]
    y = 64
    for side, name, ar, tr, en in lines:
        bubble(s, B, y, side, name, ar, tr, en, h=78, w=600)
        y += 84
    legend(s, [(BLUE, LBLUE, "blue = emphatic tool")], y=484)


@slide(4)
def p4_work(s, B, ctx):
    frame(s, "At Work: A Team Meeting", stop=5, page=ctx.page)
    tb(s.shapes, X0, 62, CW, 22, Para("WORKPLACE · The team leader speaks "
       "(polite, semi-formal)", size=12.5, color=MUTED, align="r"),
       margins=(0, 0, 0, 0))
    data = [
        ("أَوَدّ أَن أَشْكُرَكُم {b:جَمِيعًا} عَلَى جُهُودِكُم.",
         "awaddu an ashkurakum jamīʿan ʿalā juhūdikum.",
         "I'd like to thank you all for your efforts."),
        ("{b:لَقَد} حَقَّقْنَا هَدَفَنَا هَذَا الشَّهْر.",
         "laqad ḥaqqaqnā hadafanā hādhā sh-shahra.",
         "We have truly achieved our goal this month."),
        ("وَ{b:قَد} أَكَّد الْمُدِير {b:نَفْسُه} {b:أَنّ} الْمَشْرُوع نَاجِح.",
         "wa-qad akkada l-mudīru nafsuhu anna l-mashrūʿa nājiḥun.",
         "The director himself has confirmed that the project is a success."),
        ("{b:إِنَّمَا} نَحْتَاج إِلَى وَقْت إِضَافِيّ قَلِيل.",
         "innamā naḥtāju ilā waqtin iḍāfiyyin qalīlin.",
         "We only need a little extra time."),
    ]
    y = 90
    for ar, tr, en in data:
        b = exbox(s.shapes, X0, y, CW, 88, ar, tr, en)
        B.click(b, "wipe")
        y += 96
    legend(s, [(BLUE, LBLUE, "blue = emphatic tool")], y=484)


@slide(4)
def p4_formal(s, B, ctx):
    frame(s, "Announcements and Formal Statements", stop=5, page=ctx.page)
    data = [
        ("AIRPORT ANNOUNCEMENT", "نُؤَكِّد {b:أَنّ} الرِّحْلَة سَتُقْلِع فِي مَوْعِدِهَا.",
         "nuʾakkidu anna r-riḥlata sa-tuqliʿu fī mawʿidihā.",
         "We confirm that the flight will depart on time."),
        ("AIRLINE STATEMENT", "{b:إِنّ} سَلَامَة الْمُسَافِرِين هِيَ أَوْلَوِيَّتُنَا الْأُولَى.",
         "inna salāmata l-musāfirīna hiya awlawiyyatunā l-ūlā.",
         "Passenger safety is our first priority."),
        ("NEWS", "وَ{b:قَد} أَعْلَنَت الْوِزَارَة {b:أَنّ} الْمَدَارِس سَتُفْتَح الْأُسْبُوع الْقَادِم.",
         "wa-qad aʿlanati l-wizāratu anna l-madārisa sa-tuftaḥu l-usbūʿa l-qādima.",
         "The ministry has announced that schools will open next week."),
        ("SPEECH", "{b:إِنَّه لَ}شَرَف لِي أَن أَقِف أَمَامَكُم الْيَوْم.",
         "innahu la-sharafun lī an aqifa amāmakumu l-yawma.",
         "It is truly an honour for me to stand before you today."),
    ]
    y = 64
    for lab, ar, tr, en in data:
        b = exbox(s.shapes, X0, y, CW, 96, ar, tr, en, label=lab, fill=TINT2)
        B.click(b, "wipe")
        y += 104
    legend(s, [(BLUE, LBLUE, "blue = emphatic tool")], y=484)


@slide(4)
def p4_ladder(s, B, ctx):
    frame(s, "Practice Ladder: From Recognition to Production", stop=5,
          page=ctx.page)
    drill_header(s, y=62, labels=("Task", "Rung", "Answer"))
    items = [
        dict(q=["لَقَد فَاز الْفَرِيق", "Find the emphatic tool."],
             cue="{k:1 · Recognise}", a=["{g:لَقَد}", "~laqad fāza l-farīqu"]),
        dict(q=["إِنَّمَا أَنَا ضَيْف", "What kind of emphasis is this?"],
             cue="{k:2 · Classify}", a="{g:✓ Restriction}\n'I'm only a guest.'"),
        dict(q=["نَجَح الطُّلَّاب ــــــ", "Complete: 'all of them'."],
             cue="{k:3 · Complete}", a=["{g:كُلُّهُم}", "~najaḥa ṭ-ṭullābu kulluhum"]),
        dict(q=["الْفِكْرَة رَائِعَة", "Rewrite with إِنّ … لَـ."],
             cue="{k:4 · Transform}",
             a=["{g:إِنّ الْفِكْرَة لَرَائِعَة}", "~inna l-fikrata la-rāʾiʿatun"]),
        dict(q="Tell a colleague: 'I have (really) finished the report.'",
             cue="{k:5 · Produce}",
             a=["{g:لَقَد أَنْهَيْت التَّقْرِير}", "~laqad anhaytu t-taqrīra"]),
    ]
    drill(s, B, items, y0=86, rowh=74, gap=7)


@slide(4)
def p4_natural(s, B, ctx):
    frame(s, "Which Sounds Natural?", stop=5, page=ctx.page)
    drill_header(s, y=62, labels=("Situation", "Option أ (A) · Option ب (B)", "Answer"),
                 widths=(30, 300, 290, 204))
    items = [
        dict(q="Describing your daily routine to a friend.",
             cue=["أ: أَسْتَيْقِظ فِي السَّابِعَة", "ب: لَقَد أَسْتَيْقِظ فِي السَّابِعَة"],
             a=["{g:✓ أ (A)}", Para("Routine needs no emphasis, and لَقَد "
                                     "needs a past verb.", size=11.5, kind="en")]),
        dict(q="Thanking everyone at the end of a meeting.",
             cue=["أ: شُكْرًا لَكُم جَمِيعًا", "ب: شُكْرًا لِجَمِيع لَكُم"],
             a=["{g:✓ أ (A)}", Para("The fixed form is جَمِيعًا.", size=11.5,
                                     kind="en")]),
        dict(q="Someone says: 'You didn't send it!'",
             cue=["أ: أَرْسَلْتُه", "ب: لَقَد أَرْسَلْتُه!"],
             a=["{g:✓ ب (B)}", Para("Emphasis answers a denial.", size=11.5,
                                     kind="en")]),
        dict(q="Saying: 'I only want to help.'",
             cue=["أ: إِنَّمَا أُرِيد أَن أُسَاعِد", "ب: إِنّ أُرِيد أَن أُسَاعِد"],
             a=["{g:✓ أ (A)}", Para("A verb cannot follow إِنّ directly.",
                                     size=11.5, kind="en")]),
        dict(q="Saying: 'We're in the same class.'",
             cue=["أ: نَحْن فِي نَفْس الصَّفّ", "ب: نَحْن فِي نَفْس صَفّ"],
             a=["{g:✓ أ (A)}", Para("'The same' needs a definite noun.",
                                     size=11.5, kind="en")]),
    ]
    drill(s, B, items, y0=86, rowh=76, gap=6, widths=(30, 300, 290, 204),
          ar_size=16)


@slide(4)
def p4_review(s, B, ctx):
    frame(s, "Cumulative Review", stop=5, page=ctx.page)
    drill_header(s, y=62, labels=("Sentence", "Hint", "Answer"))
    B_ = "ــــــ"
    items = [
        dict(q=[f"{B_} الْوَقْت كَالسَّيْف", "Time is (truly) like a sword. (proverb)"],
             cue="start of a nominal sentence",
             a=["{g:إِنّ}", "~inna l-waqta ka-s-sayfi"]),
        dict(q=[f"أَعْرِف {B_} مَشْغُول", "I know (that) you are busy."],
             cue="after 'I know' + 'you'",
             a=["{g:أَنَّك}", "~aʿrifu annaka mashghūlun"]),
        dict(q=[f"{B_} وَصَل الطَّرْد", "The parcel has arrived."],
             cue="past verb, completion",
             a=["{g:قَد} / {g:لَقَد}", "~qad waṣala ṭ-ṭardu"]),
        dict(q=[f"{B_} يَتَغَيَّر الْمَوْعِد", "The appointment may change."],
             cue="present verb, 'may'",
             a=["{g:قَد}", "~qad yataghayyaru l-mawʿidu"]),
        dict(q=[f"قَابَلْت الْكَاتِب {B_}", "I met the writer himself."],
             cue="'himself' (object)",
             a=["{g:نَفْسَه}", "~qābaltu l-kātiba nafsahu"]),
        dict(q=[f"{B_} أَنَا أَسْأَل", "I'm only asking."],
             cue="'only'",
             a=["{g:إِنَّمَا}", "~innamā ana asʾalu"]),
    ]
    drill(s, B, items, y0=86, rowh=64, gap=5)


@slide(4)
def p4_personal(s, B, ctx):
    frame(s, "Your Turn: Personal Application", stop=5, page=ctx.page)
    banner(s, B, "Write one sentence for each prompt about YOUR life. Then "
           "compare with the model.")
    data = [
        ("1 · Something You Have Truly Achieved", "use لَقَد",
         "{b:لَقَد} تَعَلَّمْت الْحُرُوف الْعَرَبِيَّة {b:كُلَّهَا}.",
         "laqad taʿallamtu l-ḥurūfa l-ʿarabiyyata kullahā.",
         "I've truly learned all the Arabic letters."),
        ("2 · Something That May Happen", "use قَد + present",
         "{b:قَد} أَزُور الْمَغْرِب الصَّيْف الْقَادِم.",
         "qad azūru l-maghriba ṣ-ṣayfa l-qādima.",
         "I may visit Morocco next summer."),
        ("3 · Something You Firmly Believe", "use إِنّ",
         "{b:إِنّ} تَعَلُّم اللُّغَات مُفِيد جِدًّا.",
         "inna taʿalluma l-lughāti mufīdun jiddan.",
         "Learning languages is truly very useful."),
    ]
    w, gap = 272, 14
    x = X1 - w
    for t, use, ar, tr, en in data:
        g = card(s.shapes, x, 106, w, 300, t, [
            Para(use.upper(), size=11, color=BLUE, bold=True, align="c",
                 kind="en", space_after=6),
            Para("Model answer:", size=11, color=MUTED, align="c", kind="en"),
            Para(ar, size=18, align="c", kind="ar", space_after=0),
            Para("{t:%s}" % tr, size=12, align="c", kind="en"),
            Para(en, size=12.5, align="c", kind="en", bold=True)],
            anchor="m", title_size=13)
        B.click(g, "wipe")
        x -= w + gap
    legend(s, [(BLUE, LBLUE, "blue = emphatic tool")], y=420)


# ===================================================================== PART 5
@slide(5)
def p5_open(s, B, ctx):
    opener(s, B, ctx, 6, "Clearing the Confusion",
           "Look-alike tools, side by side",
           "الْفُرُوق الدَّقِيقَة", "al-furūq ad-daqīqa",
           "Separate the pairs learners confuse most, check any sentence in "
           "three quick steps, and prove it in a final drill.",
           [("that vs. indeed", "إِنّ / أَنّ", "inna / anna"),
            ("strong vs. plain", "لَقَد / قَد", "laqad / qad"),
            ("only vs. indeed", "إِنَّمَا / إِنّ", "innamā / inna"),
            ("-self vs. very same", "نَفْس / عَيْن", "nafs / ʿayn"),
            ("all vs. all", "كُلّ / جَمِيع", "kull / jamīʿ")],
           ["Explain the difference inside each look-alike pair.",
            "Use three diagnostic checks to test any sentence.",
            "Correct errors and complete the final mastery drill."])


def pair_slide(s, B, ctx, title, right, left, note_txt, note_label="KEY POINT"):
    frame(s, title, stop=6, page=ctx.page)
    w = 416
    a = card(s.shapes, X1 - w, 66, w, 300, right[0], right[1], title_size=15)
    b = card(s.shapes, X0, 66, w, 300, left[0], left[1], title_size=15,
             color=BLUE, body_fill=LBLUE)
    B.click(a, "wipe")
    B.click(b, "wipe")
    n = note(s.shapes, X0, 380, CW, 100, note_txt, label=note_label, size=14)
    B.click(n)


def exp(ar, tr, en):
    return [Para(ar, size=19, align="r", kind="ar", space_after=0),
            Para("{t:%s}  —  %s" % (tr, en), size=12, align="r", kind="en",
                 space_after=8)]


@slide(5)
def p5_inna_anna(s, B, ctx):
    pair_slide(s, B, ctx, "Confusion 1: إِنّ vs. أَنّ",
               ("Opens the Statement: inna · إِنّ", [
                   Para("start of sentence · after قَال", size=13, color=MUTED,
                        align="c", space_after=6)]
                + exp("{b:إِنّ} الدَّرْس سَهْل", "inna d-darsa sahlun",
                      "The lesson really is easy.")
                + exp("قَال {b:إِنّ} الدَّرْس سَهْل", "qāla inna d-darsa sahlun",
                      "He said the lesson is easy.")),
               ("'That', Inside: anna · أَنّ", [
                   Para("after verbs of knowing · after prepositions",
                        size=13, color=MUTED, align="c", space_after=6)]
                + exp("عَرَفْت {b:أَنّ} الدَّرْس سَهْل",
                      "ʿaraftu anna d-darsa sahlun", "I found out the lesson is easy.")
                + exp("{b:لِأَنّ} الدَّرْس سَهْل", "li-anna d-darsa sahlun",
                      "because the lesson is easy")),
               "Four look-alikes: إِنّ {t:inna} 'indeed' · أَنّ {t:anna} 'that' "
               "· إِن {t:in} 'if' · أَن {t:an} 'to' (+ verb). Shadda = the "
               "emphatic pair.", note_label="FOUR LOOK-ALIKES")


@slide(5)
def p5_qad(s, B, ctx):
    pair_slide(s, B, ctx, "Confusion 2: قَد with Past vs. Present",
               ("With a PAST Verb: Completion", [
                   Para("'already', 'has done'", size=13, color=MUTED,
                        align="c", space_after=6)]
                + exp("{b:قَد} نَسِيت كَلِمَة السِّرّ", "qad nasītu kalimata s-sirri",
                      "I have forgotten the password.")
                + exp("{b:قَد} بَدَأ الْفِيلْم", "qad badaʾa l-fīlmu",
                      "The film has started.")),
               ("With a PRESENT Verb: Possibility", [
                   Para("'may', 'might', 'sometimes'", size=13, color=MUTED,
                        align="c", space_after=6)]
                + exp("{b:قَد} أَنْسَى كَلِمَة السِّرّ", "qad ansā kalimata s-sirri",
                      "I may forget the password.")
                + exp("{b:قَد} يَبْدَأ الْفِيلْم مُتَأَخِّرًا",
                      "qad yabdaʾu l-fīlmu mutaʾakhkhiran",
                      "The film may start late.")),
               "Cover قَد and read the verb. Past verb = 'has done'. "
               "Present verb = 'may do'.", note_label="QUICK TEST")


@slide(5)
def p5_laqad_qad(s, B, ctx):
    frame(s, "Confusion 3: لَقَد vs. قَد", stop=6, page=ctx.page)
    G = lambda t: ("{g:✓} " + t, LGREEN)
    R = lambda t: ("{r:✗} " + t, LRED)
    rows = [
        ("before a past verb", G("قَد وَصَل"), G("لَقَد وَصَل")),
        ("before a present verb", G("قَد يَصِل\n'may arrive'"), R("not used")),
        ("strength", "plain, factual", "strong, emotional"),
        ("typical use", "FORMAL reports: وَقَد …", "news, surprise, feelings, replies"),
        ("with negation", G("قَد لَا يَصِل"), R("not used")),
    ]
    table(s, B, X0, 62, [284, 280, 280], ["Feature", "قَد (qad)",
          "لَقَد (laqad)"], rows, rowh=56, ar_size=18, en_size=12.5)
    n = note(s.shapes, X0, 376, CW, 100, "Both + past verb = completion. "
             "Choose لَقَد when the moment deserves force: لَقَد نَجَحْت! "
             "({t:laqad najaḥtu!}, 'I actually passed!').", label="KEY POINT",
             size=14)
    B.click(n)
    legend(s, [(GREEN, LGREEN, "✓ green = used"), (RED, LRED, "✗ red = not used")],
           y=484)


@slide(5)
def p5_innama(s, B, ctx):
    pair_slide(s, B, ctx, "Confusion 4: إِنَّمَا vs. Ordinary Emphasis",
               ("Strengthens: inna · إِنّ", [
                   Para("the statement is TRUE", size=13, color=MUTED,
                        align="c", space_after=6)]
                + exp("{b:إِنّ} أَحْمَد طَالِب", "inna Aḥmada ṭālibun",
                      "Ahmad is indeed a student.")
                + exp("{b:إِنّ} الْمُشْكِلَة بَسِيطَة",
                      "inna l-mushkilata basīṭatun", "The problem really is simple.")),
               ("Limits: innamā · إِنَّمَا", [
                   Para("the statement is ONLY this", size=13, color=MUTED,
                        align="c", space_after=6)]
                + exp("{b:إِنَّمَا} أَحْمَد طَالِب", "innamā Aḥmadu ṭālibun",
                      "Ahmad is only a student.")
                + exp("{b:إِنَّمَا} الْمُشْكِلَة فِي الْوَقْت",
                      "innamā l-mushkilatu fī l-waqti",
                      "The problem is only the time.")),
               "Don't confuse إِنَّمَا (one word) with إِنّ مَا (two words, "
               "'indeed what'): "
               "إِنّ مَا قُلْتَه صَحِيح ({t:inna mā qultahu ṣaḥīḥun}) — "
               "'What you said is correct.'", note_label="WRITTEN APART = DIFFERENT")


@slide(5)
def p5_nafs_ayn(s, B, ctx):
    pair_slide(s, B, ctx, "Confusion 5: نَفْس vs. عَيْن",
               ("Everyday and Formal: nafs · نَفْس", [
                   Para("'-self' after a noun · 'the same' before a noun · "
                        "'myself': بِنَفْسِي", size=12.5, color=MUTED,
                        align="c", space_after=6)]
                + exp("جَاءَت الْمُدِيرَة {b:نَفْسُهَا}", "jāʾati l-mudīratu nafsuhā",
                      "The manager herself came.")
                + exp("{b:نَفْس} الْمُشْكِلَة", "nafsu l-mushkilati",
                      "the same problem")),
               ("Formal Only: ʿayn · عَيْن", [
                   Para("'the very same' after a noun · also 'eye', 'spring'",
                        size=12.5, color=MUTED, align="c", space_after=6)]
                + exp("هِيَ الْمُشْكِلَة {b:عَيْنُهَا}", "hiya l-mushkilatu ʿaynuhā",
                      "It is the very same problem.")
                + exp("رَأَيْتُه بِعَيْنَيّ", "raʾaytuhu bi-ʿaynayya",
                      "I saw it with my own eyes. (literal 'eyes')")),
               "Both follow the noun's case and take a matching suffix. "
               "Only نَفْس works as 'myself / yourself': بِنَفْسِي. When in "
               "doubt, use نَفْس.", note_label="KEY POINT")


@slide(5)
def p5_kull_jami(s, B, ctx):
    pair_slide(s, B, ctx, "Confusion 6: كُلّ vs. جَمِيع",
               ("The All-Rounder: kull · كُلّ", [
                   Para("every · all · the whole · emphasis",
                        size=13, color=MUTED, align="c", space_after=6)]
                + exp("{b:كُلّ} صَبَاح", "kulla ṣabāḥin", "every morning")
                + exp("الْبَيْت {b:كُلُّه} نَظِيف", "al-baytu kulluhu naẓīfun",
                      "The whole house is clean.")),
               ("'All' of a Group: jamīʿ · جَمِيع", [
                   Para("all + definite plural · emphasis · all together",
                        size=13, color=MUTED, align="c", space_after=6)]
                + exp("{b:جَمِيع} الْغُرَف نَظِيفَة",
                      "jamīʿu l-ghurafi naẓīfatun", "All the rooms are clean.")
                + exp("ادْخُلُوا {b:جَمِيعًا}", "udkhulū jamīʿan",
                      "Come in, all of you.")),
               "For 'every': كُلّ only. For 'the whole': كُلّ only. For 'all "
               "together': جَمِيعًا. For 'all the …': either.",
               note_label="QUICK TEST")


@slide(5)
def p5_exclam(s, B, ctx):
    pair_slide(s, B, ctx, "Confusion 7: Emphasis vs. Exclamation",
               ("Emphasis — Asserts a Fact", [
                   Para("ends with a full stop · answers doubt", size=13,
                        color=MUTED, align="c", space_after=6)]
                + exp("{b:إِنّ} الْبَحْر جَمِيل.", "inna l-baḥra jamīlun.",
                      "The sea really is beautiful.")
                + exp("{b:لَقَد} كَان الْعَرْض رَائِعًا.",
                      "laqad kāna l-ʿarḍu rāʾiʿan.", "The show was truly great.")),
               ("Exclamation — Expresses Wonder", [
                   Para("ends with ! · Chapter 2.12", size=13, color=MUTED,
                        align="c", space_after=6)]
                + exp("مَا أَجْمَل الْبَحْر!", "mā ajmala l-baḥra!",
                      "How beautiful the sea is!")
                + exp("يَا لَه مِن عَرْض رَائِع!", "yā lahu min ʿarḍin rāʾiʿin!",
                      "What a wonderful show!")),
               "Emphasis says 'this IS true'. Exclamation says 'I'm amazed "
               "by this'. A sentence can be both true and amazing, but the "
               "tools are different.", note_label="KEY POINT")


@slide(5)
def p5_affirm(s, B, ctx):
    frame(s, "Confusion 8: Emphasis vs. Affirmation and Negation", stop=6,
          page=ctx.page)
    w, gap = 272, 14
    x = X1 - w
    data = [
        ("Affirmation = Reply Words", [
            Para("They answer; they are not emphatic tools.", size=12.5,
                 space_after=6)]
         + exp("نَعَم", "naʿam", "yes")
         + exp("أَجَل", "ajal", "yes (FORMAL)")
         + exp("— أَلَم تَفْهَم؟ — بَلَى.", "a-lam tafham? — balā.",
               "Didn't you understand? — Yes, I did.")),
        ("Negation Can Be Emphasized", [
            Para("Emphasis tools also work with negative statements.",
                 size=12.5, space_after=6)]
         + exp("{b:إِنّ} الْمَتْحَف لَيْس مَفْتُوحًا الْيَوْم",
               "inna l-matḥafa laysa maftūḥan l-yawma",
               "The museum really isn't open today.")
         + exp("{b:قَد} لَا يَأْتِي", "qad lā yaʾtī", "He may not come.")),
        ("Strong Denial", [
            Para("The word كَلَّا is a FORMAL, emphatic 'no' — a reply "
                 "word, not one of the ten tools.", size=12.5, space_after=6)]
         + exp("كَلَّا", "kallā", "No, not at all!")
         + exp("لَا، أَبَدًا", "lā, abadan", "No, never! (EVERYDAY)")),
    ]
    for t, body in data:
        g = card(s.shapes, x, 66, w, 340, t, body, title_size=14)
        B.click(g, "wipe")
        x -= w + gap
    n = note(s.shapes, X0, 418, CW, 62, "After a NEGATIVE question, 'yes' is "
             "بَلَى. Answering نَعَم there would mean 'right, I didn't'.",
             label="WATCH OUT", size=14)
    B.click(n)


@slide(5)
def p5_diag(s, B, ctx):
    frame(s, "Three Quick Diagnostic Checks", stop=6, page=ctx.page)
    data = [
        ("CHECK 1 · What Comes Next?", [
            Para("Noun or pronoun next: إِنّ · أَنّ · إِنَّمَا", size=14,
                 space_after=8),
            Para("Past verb next: قَد · لَقَد", size=14, space_after=8),
            Para("Present verb next: قَد 'may' · إِنَّمَا", size=14,
                 space_after=8),
            Para("Verb after 'want / must': أَن, not أَنّ", size=14)]),
        ("CHECK 2 · Where Is It?", [
            Para("Sentence start, or after قَال: إِنّ", size=14,
                 space_after=8),
            Para("After a verb of knowing or a preposition: أَنّ", size=14,
                 space_after=8),
            Para("After a noun, with a suffix: نَفْس · عَيْن · كُلّ · جَمِيع",
                 size=14, space_after=8),
            Para("Joined to a word: لَـ", size=14)]),
        ("CHECK 3 · What Is the Job?", [
            Para("Strengthen: إِنّ · إِنّ … لَـ", size=14, space_after=8),
            Para("Completed or possible: قَد · لَقَد", size=14, space_after=8),
            Para("Limit: إِنَّمَا", size=14, space_after=8),
            Para("Identity: نَفْس · عَيْن", size=14, space_after=8),
            Para("Totality: كُلّ · جَمِيع", size=14)]),
    ]
    w, gap = 272, 14
    x = X1 - w
    for t, body in data:
        g = card(s.shapes, x, 66, w, 330, t, body, title_size=14.5)
        B.click(g, "wipe")
        x -= w + gap
    n = note(s.shapes, X0, 410, CW, 70, "Run the checks in order, starting "
             "from the right-hand card: next word, then position, then job.",
             label="HOW TO USE", size=14)
    B.click(n)


@slide(5)
def p5_correct(s, B, ctx):
    frame(s, "Correction Exercise: Find and Fix", stop=6, page=ctx.page)
    drill_header(s, y=62, labels=("Sentence with an error", "Hint",
                 "Corrected"), widths=(30, 290, 210, 294))
    items = [
        dict(q=["إِنّ الطَّقْس بَارِدًا", "The weather is indeed cold."],
             cue="predicate ending",
             a=["{g:إِنّ الطَّقْس بَارِد}", "~inna ṭ-ṭaqsa bāridun"]),
        dict(q=["لَقَد أَدْرُس الْعَرَبِيَّة كُلّ يَوْم", "I study Arabic every day."],
             cue="verb form + routine",
             a=["{g:أَدْرُس الْعَرَبِيَّة كُلّ يَوْم}",
                "~adrusu l-ʿarabiyyata kulla yawmin"]),
        dict(q=["أَظُنّ إِنّ الْفِيلْم طَوِيل", "I think the film is long."],
             cue="after a verb of thinking",
             a=["{g:أَظُنّ أَنّ الْفِيلْم طَوِيل}",
                "~aẓunnu anna l-fīlma ṭawīlun"]),
        dict(q=["قَابَلْت الْمُدِيرَة نَفْسَه", "I met the manager herself."],
             cue="suffix agreement",
             a=["{g:قَابَلْت الْمُدِيرَة نَفْسَهَا}",
                "~qābaltu l-mudīrata nafsahā"]),
        dict(q=["جَمِيع يَوْم أَشْرَب الْقَهْوَة", "Every day I drink coffee."],
             cue="'every'",
             a=["{g:كُلّ يَوْم أَشْرَب الْقَهْوَة}",
                "~kulla yawmin ashrabu l-qahwata"]),
        dict(q=["قَد سَيَصِل الْقِطَار", "The train may arrive."],
             cue="future after قَد?",
             a=["{g:قَد يَصِل الْقِطَار}", "~qad yaṣilu l-qiṭāru"]),
    ]
    drill(s, B, items, y0=86, rowh=64, gap=5, widths=(30, 290, 210, 294))


@slide(5)
def p5_mastery(s, B, ctx):
    frame(s, "Final Mastery Drill", stop=6, page=ctx.page)
    tb(s.shapes, X0, 62, CW, 24, Para("Complete each sentence. Answers on the "
       "next slide.", size=13, color=MUTED, align="r"), margins=(0, 0, 0, 0))
    B_ = "ــــــ"
    items = [
        (f"{B_} الصِّدْق مَنْجَاة", "Honesty is truly salvation. (proverb)"),
        (f"سَمِعْت {B_} الْمَتْجَر سَيُغْلَق مُبَكِّرًا", "I heard that the shop will close early."),
        (f"{B_} وَصَلْنَا أَخِيرًا!", "We've finally arrived!"),
        (f"{B_} يَتَأَخَّر الرَّدّ قَلِيلًا", "The reply may be a little late."),
        (f"إِنّ هَذَا {B_}خَبَر سَعِيد", "This is truly happy news."),
        (f"{B_} أَنَا أُعَبِّر عَن رَأْيِي", "I'm only expressing my opinion."),
        (f"اسْتَقْبَلَنَا الرَّئِيس {B_}", "The president himself received us."),
        (f"أُحِبّ هَذِه الْأُغْنِيَات {B_}", "I love all these songs."),
    ]
    w, h, gx, gy = 416, 90, 12, 10
    for i, (ar, en) in enumerate(items):
        col, row = i % 2, i // 2
        x = X1 - w - col * (w + gx)
        y = 92 + row * (h + gy)
        g = group(s.shapes)
        circ = shape(g.shapes, MSO_SHAPE.OVAL, x + w - 34, y + h / 2 - 13, 26,
                     26, WINE)
        write_paras(circ.text_frame, [Para(str(i + 1), size=12, color=WHITE,
                    bold=True, align="c")], "m", (0, 0, 0, 0))
        tb(g.shapes, x, y, w - 40, h, [Para(ar, size=19, align="r", kind="ar"),
           Para(en, size=12, align="r", kind="en")], fill=WHITE, line=LINE,
           radius=8, anchor="m")
        B.click(g, "wipe")


@slide(5)
def p5_mastery_ans(s, B, ctx):
    frame(s, "Final Mastery Drill — Answers", stop=6, page=ctx.page)
    rows = [
        ("1", "{g:إِنّ} الصِّدْق مَنْجَاة\n~inna ṣ-ṣidqa manjātun", "sentence start"),
        ("2", "سَمِعْت {g:أَنّ} الْمَتْجَر سَيُغْلَق مُبَكِّرًا\n~samiʿtu anna l-matjara sa-yughlaqu mubakkiran", "after 'I heard'"),
        ("3", "{g:لَقَد} وَصَلْنَا أَخِيرًا!\n~laqad waṣalnā akhīran!", "past verb + feeling"),
        ("4", "{g:قَد} يَتَأَخَّر الرَّدّ قَلِيلًا\n~qad yataʾakhkharu r-raddu qalīlan", "present verb = 'may'"),
        ("5", "إِنّ هَذَا {g:لَ}خَبَر سَعِيد\n~inna hādhā la-khabarun saʿīdun", "لَـ on the predicate"),
        ("6", "{g:إِنَّمَا} أَنَا أُعَبِّر عَن رَأْيِي\n~innamā ana uʿabbiru ʿan raʾyī", "restriction"),
        ("7", "اسْتَقْبَلَنَا الرَّئِيس {g:نَفْسُه}\n~istaqbalanā r-raʾīsu nafsuhu", "identity; subject takes {t:-u}"),
        ("8", "أُحِبّ هَذِه الْأُغْنِيَات {g:كُلَّهَا}\n~uḥibbu hādhihi l-ughniyāti kullahā", "things take {t:-hā}"),
    ]
    rows = [(Para(a, size=13, color=WINE, bold=True, align="c"), b, c_)
            for a, b, c_ in rows]
    half = 4
    table(s, B, X0 + 432, 62, [40, 372], ["#", "Answer"],
          [r[:2] for r in rows[:half]], rowh=88, ar_size=17, en_size=11.5)
    table(s, B, X0, 62, [40, 372], ["#", "Answer"],
          [r[:2] for r in rows[half:]], rowh=88, ar_size=17, en_size=11.5)
    tb(s.shapes, X0, 446, CW, 34, Para("Reasons — 1 sentence start · 2 after "
       "'I heard' · 3 past verb + feeling · 4 present verb = 'may' · 5 لَـ on "
       "the predicate · 6 restriction · 7 '-self', subject · 8 things take {t:-hā}",
       size=11, color=MUTED, align="c"), anchor="m", margins=(0, 0, 0, 0))
    legend(s, [(GREEN, LGREEN, "green = the answer")], y=482)


# ===================================================================== CLOSE
@slide("close")
def s_summary(s, B, ctx):
    frame(s, "Chapter Summary: Emphatic Tools", page=ctx.page)
    rows = [
        ("Particles on a sentence", "إِنّ · أَنّ · إِنَّمَا",
         "strengthen · report 'that' · limit"),
        ("Particles on a verb", "قَد · لَقَد", "completion (past) · possibility (present)"),
        ("Attached prefix", "لَـ", "reinforces one word; إِنّ … لَـ"),
        ("Nouns for emphasis", "نَفْس · عَيْن · كُلّ · جَمِيع",
         "identity · totality — with a matching suffix"),
        ("No tool", "repetition · stress · context", "emphasis in speech"),
    ]
    table(s, B, X0, 62, [230, 280, 334], ["Kind", "Tools", "Job"], rows,
          rowh=56, ar_size=19, en_size=12.5)
    n = tb(s.shapes, X0, 390, CW, 96, [
        Para("Next stop:", size=13, color=MUTED, align="c"),
        Para("Chapter 2.11 · Vocative Tools", size=24, color=WINE, bold=True,
             align="c"),
        Para("أَدَوَات النِّدَاء  {t:· adawāt an-nidāʾ}", size=19, align="c",
             kind="ar")], fill=TINT, line=WINE, radius=12, anchor="m")
    B.click(n, "wipe")


# =================================================================== BUILD
def main(out):
    ctx = Ctx()
    # pass 1: slide ranges per part (1-based page numbers)
    for i, (part, _) in enumerate(SLIDES, 1):
        a, b = ctx.ranges.get(part, (i, i))
        ctx.ranges[part] = (min(a, i), max(b, i))
    prs = new_presentation()
    blank = prs.slide_layouts[6]
    for i, (part, fn) in enumerate(SLIDES, 1):
        ctx.page = i
        s = prs.slides.add_slide(blank)
        B = Builds()
        fn(s, B, ctx)
        attach_timing(s, B)
    prs.save(out)
    print(f"saved {out}: {len(SLIDES)} slides")
    for k in [0, 1, 2, 3, 4, 5]:
        print(f"  Part {k} (Stop {k + 1}): slides {ctx.ranges[k][0]}–{ctx.ranges[k][1]}")
    return ctx


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "out/Chapter-2.10-Emphatic-Tools.pptx"
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    main(out)
