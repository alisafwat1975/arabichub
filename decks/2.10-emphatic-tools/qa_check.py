"""Automated checks on the built deck (run after build_ch2_10.py).

Checks: final-letter convention, Latin ? , ; inside RTL paragraphs, slide
titles that begin with Arabic, leftover placeholders, opener slide ranges,
and build animations present.
"""
import re
import sys
import unicodedata
from pptx import Presentation
from pptx.oxml.ns import qn

MARKS = set(chr(c) for c in range(0x064B, 0x0653)) | {"ٰ"}
SHADDA = "ّ"
AR = re.compile(r"[؀-ۿ]")


def final_mark_errors(text):
    errs = []
    for i, ch in enumerate(text):
        if ch in MARKS and ch != SHADDA:
            j = i + 1
            while j < len(text) and text[j] in MARKS:
                j += 1
            nxt = text[j] if j < len(text) else " "
            if unicodedata.category(nxt) not in ("Lo", "Lm"):
                errs.append(text[max(0, i - 8):i + 2])
    return errs


def para_text(p):
    return "".join(t.text or "" for t in p.iter(qn("a:t")))


def main(path):
    prs = Presentation(path)
    problems = []
    n = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        sld = slide._element
        title_seen = False
        for p in sld.iter(qn("a:p")):
            txt = para_text(p).replace("‏", "").replace("‎", "")
            if not txt.strip():
                continue
            ppr = p.find(qn("a:pPr"))
            rtl = ppr is not None and ppr.get("rtl") == "1"
            for e in final_mark_errors(txt):
                problems.append(f"slide {i}: vowel on final letter near «{e}»")
            if rtl and re.search(r"[?,;]", txt):
                problems.append(f"slide {i}: Latin punctuation in RTL «{txt[:40]}»")
            if re.search(r"TODO|PLACEHOLDER|lorem|slides x|x–y", txt, re.I):
                problems.append(f"slide {i}: placeholder text «{txt[:40]}»")
            if "→" in txt:
                problems.append(f"slide {i}: right-pointing arrow «{txt[:40]}»")
        # the slide title is the first sizeable wine text box after sidebar
        texts = [para_text(p) for p in sld.iter(qn("a:p"))]
        for t in texts:
            if t.strip() in ("2.10", "Emphatic Tools", "2.11", "2.12",
                             "Vocative Tools", "Exclamation Tools"):
                continue
            first = t.replace("‏", "").replace("‎", "").strip()
            if first and AR.match(first[0]) and i != 1:
                problems.append(f"slide {i}: title starts with Arabic «{first[:30]}»")
            break
        if i > 1 and sld.find(qn("p:timing")) is None and i != n:
            problems.append(f"slide {i}: no build animation")
        # opener footer ranges
        for t in texts:
            m = re.match(r"Stop (\d) of 6 · slides (\d+)–(\d+)", t)
            if m:
                a, b = int(m.group(2)), int(m.group(3))
                if not (a == i and a <= b <= n):
                    problems.append(f"slide {i}: bad opener range {t}")
    print(f"{n} slides checked")
    for p in problems:
        print("  ", p)
    print("OK" if not problems else f"{len(problems)} problem(s)")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else
                  "out/Chapter-2.10-Emphatic-Tools.pptx"))
